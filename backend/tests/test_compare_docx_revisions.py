"""Stage A1 — a DOCX with unaccepted tracked changes must be compared in its *accepted*
(final) state. LibreOffice/Gotenberg renders tracked deletions as visible struck text,
which made the pixel overlay compare against text the reviewer already deleted.
"""
import zipfile

import docx
from lxml import etree

from app.services.docx_revisions import accept_all_revisions, has_revisions

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def _q(tag):
    return "{%s}%s" % (W, tag)


def _make_docx(path):
    """Body: 'Keep [ins:Added ][del:Removed ]tail.'; a 3-row table whose middle row is a
    tracked row deletion; a footer with a tracked deletion; trackRevisions on."""
    d = docx.Document()
    p = d.add_paragraph("Keep ")
    body_p = p._p
    ins = etree.SubElement(body_p, _q("ins"), {_q("id"): "1", _q("author"): "a", _q("date"): "2026-01-01T00:00:00Z"})
    r = etree.SubElement(ins, _q("r")); t = etree.SubElement(r, _q("t")); t.text = "Added "
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    dl = etree.SubElement(body_p, _q("del"), {_q("id"): "2", _q("author"): "a", _q("date"): "2026-01-01T00:00:00Z"})
    r = etree.SubElement(dl, _q("r")); dt = etree.SubElement(r, _q("delText")); dt.text = "Removed "
    dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r = etree.SubElement(body_p, _q("r")); t = etree.SubElement(r, _q("t")); t.text = "tail."

    tbl = d.add_table(rows=3, cols=2)
    for i, row in enumerate(tbl.rows):
        row.cells[0].text = f"row{i}"
        row.cells[1].text = f"val{i}"
    tr = tbl.rows[1]._tr
    trPr = tr.find(_q("trPr"))
    if trPr is None:
        trPr = etree.Element(_q("trPr")); tr.insert(0, trPr)
    etree.SubElement(trPr, _q("del"), {_q("id"): "3", _q("author"): "a", _q("date"): "2026-01-01T00:00:00Z"})

    # formatting-only revision must not leak text either
    rpr_change_p = d.add_paragraph("Styled")._p
    run = rpr_change_p.find(_q("r"))
    rPr = run.find(_q("rPr"))
    if rPr is None:
        rPr = etree.Element(_q("rPr")); run.insert(0, rPr)
    ch = etree.SubElement(rPr, _q("rPrChange"), {_q("id"): "4", _q("author"): "a", _q("date"): "2026-01-01T00:00:00Z"})
    etree.SubElement(ch, _q("rPr"))

    footer = d.sections[0].footer.paragraphs[0]
    footer.text = "Footer keep "
    fdel = etree.SubElement(footer._p, _q("del"), {_q("id"): "5", _q("author"): "a", _q("date"): "2026-01-01T00:00:00Z"})
    r = etree.SubElement(fdel, _q("r")); dt = etree.SubElement(r, _q("delText")); dt.text = "footer-gone"

    settings = d.settings.element
    etree.SubElement(settings, _q("trackRevisions"))
    d.save(path)


def _all_text(path):
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(z.read(name))
                out += [n.text or "" for n in root.iter(_q("t"), _q("delText"))]
    return " ".join(out)


def _count(path, part, tag):
    with zipfile.ZipFile(path) as z:
        return len(etree.fromstring(z.read(part)).findall(".//w:" + tag, NS))


def test_accept_all_revisions_keeps_insertions_drops_deletions(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    assert has_revisions(str(src))

    accept_all_revisions(str(src), str(dst))

    text = _all_text(dst)
    assert "Added" in text
    assert "Removed" not in text
    assert "footer-gone" not in text
    assert "Footer keep" in text
    assert not has_revisions(str(dst))
    for tag in ("ins", "del", "delText", "moveFrom", "moveTo", "rPrChange"):
        assert _count(dst, "word/document.xml", tag) == 0, tag
    assert _count(dst, "word/settings.xml", "trackRevisions") == 0


def test_accept_all_revisions_removes_deleted_table_rows(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    accept_all_revisions(str(src), str(dst))
    t = docx.Document(str(dst)).tables[0]
    assert [r.cells[0].text for r in t.rows] == ["row0", "row2"]


def test_accepted_paragraph_text_reads_as_final(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    _make_docx(src)
    accept_all_revisions(str(src), str(dst))
    first = docx.Document(str(dst)).paragraphs[0].text
    assert first == "Keep Added tail."


def test_document_without_revisions_is_untouched_in_content(tmp_path):
    src, dst = tmp_path / "plain.docx", tmp_path / "out.docx"
    d = docx.Document(); d.add_paragraph("Nothing tracked here."); d.save(src)
    assert not has_revisions(str(src))
    accept_all_revisions(str(src), str(dst))
    assert docx.Document(str(dst)).paragraphs[0].text == "Nothing tracked here."


def test_to_pdf_converts_the_accepted_docx(tmp_path, monkeypatch):
    """The Gotenberg conversion must receive the accepted copy, never the raw upload."""
    from app.services import pdf_render_service, gotenberg_client

    src = tmp_path / "in.docx"
    _make_docx(src)
    seen = {}

    def fake_convert(path, **kw):
        seen["text"] = _all_text(path)
        seen["path"] = path
        return b"%PDF-1.4 fake"

    monkeypatch.setattr(gotenberg_client, "convert_to_pdf", fake_convert)
    out = pdf_render_service.to_pdf(str(src), "docx", str(tmp_path / "out"), "old")
    assert out.endswith("old.pdf")
    assert "Removed" not in seen["text"] and "Added" in seen["text"]
    assert seen["path"] != str(src)
