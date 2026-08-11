"""DOCX -> HTML is the import seam. Structure must survive: headings stay
headings, lists stay lists, tables stay tables. Anything that degrades to a
flat paragraph here is lost to the editor permanently."""
import io

import pytest
from docx import Document

from app.services.lexical_import import LexicalImportError, docx_to_html


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_heading_becomes_a_heading_element():
    data = _docx(lambda d: d.add_heading("Charges", level=2))
    html = docx_to_html(data)
    assert "<h2>" in html and "Charges" in html


def test_paragraph_becomes_a_paragraph():
    data = _docx(lambda d: d.add_paragraph("Returns are not guaranteed."))
    assert "<p>" in docx_to_html(data)


def test_bullet_list_becomes_a_list():
    def build(d):
        d.add_paragraph("first", style="List Bullet")
        d.add_paragraph("second", style="List Bullet")

    html = docx_to_html(_docx(build))
    assert "<ul>" in html and html.count("<li>") == 2


def test_table_becomes_a_table():
    def build(d):
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "a"
        t.rows[0].cells[1].text = "b"

    html = docx_to_html(_docx(build))
    assert "<table>" in html and "<td>" in html


def test_unreadable_input_raises_rather_than_returning_empty():
    with pytest.raises(LexicalImportError):
        docx_to_html(b"this is not a docx")


# --- headers, footers and text boxes ---------------------------------------
#
# mammoth reads the document body only, but preprocessing_service extracts these
# regions too, so the compliance engine grades text the editor would otherwise
# not contain. A finding quoting a footer could not be located or corrected, and
# clean.docx — generated from the editor's HTML — would drop the disclaimer from
# the approved artifact.

def _with_footer(footer_text: str, header_text: str = "") -> bytes:
    doc = Document()
    doc.add_paragraph("Body copy.")
    doc.sections[0].footer.paragraphs[0].text = footer_text
    if header_text:
        doc.sections[0].header.paragraphs[0].text = header_text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_footer_text_reaches_the_editor():
    html = docx_to_html(_with_footer("Past performance is not indicative of future performance."))
    assert "Past performance is not indicative of future performance." in html
    assert "Page footer" in html


def test_header_text_reaches_the_editor():
    html = docx_to_html(_with_footer("f", header_text="Bajaj Life Smart Wealth Goal VII"))
    assert "Bajaj Life Smart Wealth Goal VII" in html
    assert "Page header" in html


def test_body_still_comes_first():
    """Peripherals are appended, so the document still reads in order."""
    html = docx_to_html(_with_footer("FOOTERTEXT"))
    assert html.index("Body copy.") < html.index("FOOTERTEXT")


def test_document_without_header_or_footer_gains_nothing():
    doc = Document()
    doc.add_paragraph("Body only.")
    buf = io.BytesIO()
    doc.save(buf)
    html = docx_to_html(buf.getvalue())
    assert "Page footer" not in html and "Page header" not in html


def test_peripheral_text_is_escaped_not_injected():
    html = docx_to_html(_with_footer("Terms & <conditions> apply"))
    assert "&amp;" in html and "&lt;conditions&gt;" in html
    assert "<conditions>" not in html


# --- text boxes -------------------------------------------------------------
#
# A text box is anchored in the body but sits outside the run tree, and it is
# where these documents keep their disclaimers. The analyser reads every one of
# them; the editor got only the ones mammoth happens to recognise, so a finding
# on a DrawingML box could not be corrected, clean.docx dropped its wording, and
# the draft redline read every line of it as a reviewer deletion.
#
# Both halves matter and pull opposite ways: import the missing ones, and do NOT
# re-import the ones mammoth already converted — a duplicated mandated
# disclaimer is its own compliance problem.

_DRAWING_NS = (
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)


def _drawing_box(text: str) -> str:
    """A DrawingML text box — mammoth walks into `w:drawing` but does not know
    `wps:txbx`, so the wording never reaches the body HTML."""
    from docx.oxml.ns import nsdecls

    return (
        f"<w:p {nsdecls('w')} {_DRAWING_NS}><w:r><w:drawing><wp:inline><a:graphic>"
        f"<a:graphicData><wps:wsp><wps:txbx><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        f"</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>"
        f"</wp:inline></w:drawing></w:r></w:p>"
    )


def _vml_box(text: str) -> str:
    """The VML flavour — `v:textbox` IS in mammoth's element map, so this one
    converts as body copy on its own."""
    from docx.oxml.ns import nsdecls

    return (
        f"<w:p {nsdecls('w')}><w:r><w:pict>"
        f'<v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
        f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        f"</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
    )


def _with_boxes(*xml: str) -> bytes:
    """python-docx cannot add a text box, so the XML Word writes is built by
    hand and appended to the body."""
    from docx.oxml import parse_xml

    doc = Document()
    doc.add_paragraph("Body copy.")
    for fragment in xml:
        doc.element.body.append(parse_xml(fragment))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_drawingml_text_box_reaches_the_editor():
    html = docx_to_html(_with_boxes(_drawing_box("Tax benefits are subject to change.")))
    assert "Tax benefits are subject to change." in html
    assert "Text box" in html


def test_a_text_box_mammoth_already_converted_is_not_imported_twice():
    html = docx_to_html(_with_boxes(_vml_box("Tax benefits are subject to change.")))
    assert html.count("Tax benefits are subject to change.") == 1
    assert "Text box" not in html, "already in the body; labelling it again duplicates it"


def test_each_unconverted_text_box_is_its_own_labelled_block():
    html = docx_to_html(
        _with_boxes(_drawing_box("First disclaimer."), _drawing_box("Second disclaimer."))
    )
    assert "Text box 1" in html and "Text box 2" in html
    assert html.index("First disclaimer.") < html.index("Second disclaimer.")


def test_repeated_text_box_wording_is_not_duplicated():
    """Same dedup rule as headers and footers: a repeat is not new wording."""
    html = docx_to_html(
        _with_boxes(_drawing_box("Same disclaimer."), _drawing_box("Same disclaimer."))
    )
    assert html.count("Same disclaimer.") == 1


def test_text_box_content_is_escaped_not_injected():
    html = docx_to_html(_with_boxes(_drawing_box("Terms &amp; &lt;conditions&gt; apply")))
    assert "&amp;" in html and "&lt;conditions&gt;" in html
    assert "<conditions>" not in html


def test_document_without_a_text_box_gains_nothing():
    assert "Text box" not in docx_to_html(_docx(lambda d: d.add_paragraph("Body only.")))
