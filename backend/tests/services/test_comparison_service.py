from pathlib import Path
from unittest.mock import patch, MagicMock

import pytest

from app.services.comparison_service import (
    split_text_paragraphs,
    extract_docx_paragraphs,
    extract_pdf_paragraphs,
    extract_paragraphs,
)


def test_split_text_paragraphs_splits_on_blank_lines():
    text = "First paragraph.\n\nSecond paragraph.\n\n\nThird paragraph."
    assert split_text_paragraphs(text) == [
        "First paragraph.",
        "Second paragraph.",
        "Third paragraph.",
    ]


def test_split_text_paragraphs_empty_string_returns_empty_list():
    assert split_text_paragraphs("") == []
    assert split_text_paragraphs("   ") == []


def test_extract_docx_paragraphs_reads_paragraphs_and_tags_headings(tmp_path: Path):
    from docx import Document
    doc = Document()
    doc.add_heading("Policy Overview", level=1)
    doc.add_paragraph("This plan offers guaranteed returns.")
    doc.add_paragraph("")  # blank paragraph should be skipped
    doc.add_paragraph("Terms and conditions apply.")
    file_path = tmp_path / "sample.docx"
    doc.save(str(file_path))

    paragraphs = extract_docx_paragraphs(str(file_path))

    assert paragraphs == [
        "## Policy Overview",
        "This plan offers guaranteed returns.",
        "Terms and conditions apply.",
    ]


def test_extract_pdf_paragraphs_splits_page_text_on_blank_lines():
    fake_page = MagicMock()
    fake_page.extract_text.return_value = "Para one.\n\nPara two."
    fake_pdf = MagicMock()
    fake_pdf.__enter__.return_value.pages = [fake_page]
    fake_pdf.__exit__.return_value = False

    with patch("pdfplumber.open", return_value=fake_pdf):
        paragraphs = extract_pdf_paragraphs("fake.pdf")

    assert paragraphs == ["Para one.", "Para two."]


def test_extract_paragraphs_dispatches_by_content_type(tmp_path: Path):
    text_paragraphs = extract_paragraphs(None, "text", "Hello world.\n\nSecond line.")
    assert text_paragraphs == ["Hello world.", "Second line."]

    from docx import Document
    doc = Document()
    doc.add_paragraph("Docx body text.")
    file_path = tmp_path / "sample.docx"
    doc.save(str(file_path))
    docx_paragraphs = extract_paragraphs(str(file_path), "docx")
    assert docx_paragraphs == ["Docx body text."]


def test_extract_paragraphs_requires_file_path_for_docx_and_pdf():
    with pytest.raises(ValueError):
        extract_paragraphs(None, "docx")
    with pytest.raises(ValueError):
        extract_paragraphs(None, "pdf")
