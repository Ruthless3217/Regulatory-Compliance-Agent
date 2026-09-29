"""Export service — changes derivation + Word report generation (no PDF render)."""
import io
from datetime import datetime, timezone
from types import SimpleNamespace

from docx import Document

from app.services.export_service import derive_changes, changes_report_docx


def _ann(change_id, note=None, tags=None):
    return SimpleNamespace(change_id=change_id, note=note, tags=tags or [])


def _text_comparison():
    """A completed text-only comparison (no render) with one of each block kind."""
    return SimpleNamespace(
        id="00000000-0000-0000-0000-000000000001",
        title="Brochure v1 vs v2",
        created_at=datetime(2026, 7, 12, 9, 30, tzinfo=timezone.utc),
        status="completed",
        render_status="skipped",
        render_result=None,
        old_file_path=None,
        new_file_path=None,
        diff_result=[
            {"type": "equal", "old_text": "Unchanged.", "new_text": "Unchanged."},
            {"type": "delete", "old_text": "Removed clause."},
            {"type": "insert", "new_text": "Added clause."},
            {
                "type": "replace",
                "old_words": [{"text": "old", "changed": True}, {"text": "kept", "changed": False}],
                "new_words": [{"text": "new", "changed": True}, {"text": "kept", "changed": False}],
            },
        ],
    )


def test_derive_changes_from_text_blocks():
    changes = derive_changes(_text_comparison())
    kinds = [c["kind"] for c in changes]
    assert kinds == ["removed", "added", "modified"]  # equal skipped
    assert changes[0]["change_id"] == "b1"
    assert changes[0]["kind"] == "removed"
    assert changes[0]["old_text"] == "Removed clause."
    assert changes[0]["new_text"] == ""
    assert changes[1]["change_id"] == "b2" and changes[1]["new_text"] == "Added clause."
    # replace keeps only changed words
    assert changes[2]["old_text"] == "old" and changes[2]["new_text"] == "new"


def test_derive_changes_prefers_completed_render():
    c = SimpleNamespace(
        render_status="completed",
        render_result={
            "changes": [
                {"id": "r0", "kind": "moved", "old": {"text": "A moved run."}, "new": {"text": "A moved run."}},
                {"id": "r1", "kind": "added", "new": {"text": "Only new."}},
            ]
        },
        diff_result=[],
    )
    changes = derive_changes(c)
    assert [x["change_id"] for x in changes] == ["r0", "r1"]
    assert changes[0]["kind"] == "moved"
    assert changes[1]["old_text"] == "" and changes[1]["new_text"] == "Only new."


def test_changes_report_docx_is_valid_and_includes_annotations():
    comparison = _text_comparison()
    annotations = [_ann("b1", note="check with legal", tags=["review", "critical"])]
    data = changes_report_docx(comparison, annotations)
    assert data[:2] == b"PK"  # docx is a zip container

    doc = Document(io.BytesIO(data))
    assert any("Brochure v1 vs v2" in p.text for p in doc.paragraphs)
    table = doc.tables[0]
    header = [c.text for c in table.rows[0].cells]
    assert header == ["#", "Type", "Section", "Original", "Revised", "Note", "Tags"]
    # 3 non-equal changes -> 3 body rows
    assert len(table.rows) == 1 + 3
    body_text = "\n".join(cell.text for row in table.rows[1:] for cell in row.cells)
    assert "Removed clause." in body_text
    assert "check with legal" in body_text
    assert "review, critical" in body_text
