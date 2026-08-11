"""GET /submissions/{id}/draft-diff must show the reviewer's edits and nothing
else.

The redline diffs the working copy against a baseline, and the working copy is
the Lexical editor's text. Deriving the baseline from the ANALYSER's extraction
of the same upload instead — a different reader of the same bytes, with its own
"[PAGE FOOTER]" labels, table pipes and text-box block — made every one of those
disagreements read as a reviewer edit: hundreds of phantom changes appeared the
moment anyone saved. So the document built here is deliberately one the two
readers disagree about, and `test_the_two_readers_really_do_disagree` fails
first if that stops being true and the rest of the file stops proving anything.
"""
import asyncio
import io
import uuid

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.api.routes import submissions as routes
from app.models.submission import Submission
from app.services import comparison_service, lexical_document_service
from app.services.lexical_import import docx_to_html, html_to_text

FOOTER = "Past performance is not indicative of future performance."
TEXTBOX = "Tax benefits are subject to change in tax laws."


_DRAWING_NS = (
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)


def _textbox(text: str):
    """A body paragraph carrying a DrawingML text box — python-docx cannot build
    one, and a text box is where these documents keep their disclaimers. This is
    the flavour mammoth does not convert, so it reaches the editor only through
    the peripheral pass (see test_lexical_import)."""
    return parse_xml(
        f"<w:p {nsdecls('w')} {_DRAWING_NS}><w:r><w:drawing><wp:inline><a:graphic>"
        f"<a:graphicData><wps:wsp><wps:txbx><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        f"</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic>"
        f"</wp:inline></w:drawing></w:r></w:p>"
    )


def _upload(tmp_path):
    """A DOCX with body copy, a table, a footer and a text box — every region
    the two readers render differently."""
    doc = Document()
    doc.add_heading("Charges", level=2)
    doc.add_paragraph("The premium allocation charge is deducted upfront.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Policy year"
    table.rows[0].cells[1].text = "Allocation charge"
    doc.element.body.append(_textbox(TEXTBOX))
    doc.sections[0].footer.paragraphs[0].text = FOOTER

    buf = io.BytesIO()
    doc.save(buf)
    path = tmp_path / "creative.docx"
    path.write_bytes(buf.getvalue())
    return path


def _submission(path, current=None):
    from app.services.preprocessing_service import ContextEngineeringService

    return Submission(
        id=uuid.uuid4(),
        title="t",
        content_type="docx",
        file_path=str(path),
        # What the analyser graded — the other reader, and the old baseline.
        original_content=asyncio.run(
            ContextEngineeringService(None)._extract_docx(str(path))
        ),
        current_content=current,
    )


class _FakeDB:
    """Enough Session for the route's single lookup."""

    def __init__(self, row):
        self._row = row

    def query(self, _model):
        return self

    def filter(self, *_exprs):
        return self

    def first(self):
        return self._row


def _diff(sub):
    return asyncio.run(routes.submission_draft_diff(str(sub.id), db=_FakeDB(sub)))


def _baseline(path):
    return html_to_text(docx_to_html(path.read_bytes()))


# --- the regression ---------------------------------------------------------


def test_the_two_readers_really_do_disagree(tmp_path):
    """Guards the fixture, not the code: if this document ever stopped
    provoking a disagreement, the tests below would pass vacuously."""
    path = _upload(tmp_path)
    analyser = _submission(path).original_content
    blocks = comparison_service.build_diff(
        comparison_service.split_text_paragraphs(analyser),
        comparison_service.split_text_paragraphs(_baseline(path)),
    )
    assert any(b.get("type") != "equal" for b in blocks)


def test_an_unedited_working_copy_shows_no_changes(tmp_path):
    """The bug, from the reviewer's seat: they save once, having changed
    nothing of substance, and the redline claims the document was rewritten."""
    path = _upload(tmp_path)
    res = _diff(_submission(path, current=_baseline(path)))

    assert res["changed"] == 0
    assert [b for b in res["blocks"] if b.get("type") != "equal"] == []
    assert res["edited"] is False


def test_one_changed_word_is_the_only_change(tmp_path):
    path = _upload(tmp_path)
    working = _baseline(path).replace("upfront", "annually")

    res = _diff(_submission(path, current=working))

    assert res["edited"] is True
    changed = [b for b in res["blocks"] if b.get("type") != "equal"]
    assert len(changed) == 1
    text = str(changed[0])
    assert "annually" in text and "upfront" in text


def test_text_box_and_footer_are_not_reported_as_deletions(tmp_path):
    """The two largest phantom sources: the analyser reads both regions, and
    the editor now does too — so neither may appear in the redline."""
    path = _upload(tmp_path)
    res = _diff(_submission(path, current=_baseline(path)))

    rendered = str(res["blocks"])
    assert res["changed"] == 0
    assert TEXTBOX in _baseline(path) and FOOTER in _baseline(path)
    assert "[PAGE FOOTER]" not in rendered and "[TEXT BOXES]" not in rendered


def test_no_working_copy_is_an_empty_redline(tmp_path):
    res = _diff(_submission(_upload(tmp_path)))
    assert res["changed"] == 0 and res["edited"] is False


# --- submissions with no editor lineage -------------------------------------


def test_pasted_text_still_diffs_against_the_original(tmp_path):
    """No upload, so no import to seed from: original_content IS the baseline."""
    sub = Submission(
        id=uuid.uuid4(),
        title="t",
        content_type="text",
        file_path=None,
        original_content="Returns are guaranteed.",
        current_content="Returns are not guaranteed.",
    )
    res = _diff(sub)
    assert res["edited"] is True and res["changed"] == 1


def test_an_unconvertible_upload_falls_back_to_the_analyser_text(tmp_path):
    """The editor degrades to extracted text for this submission, so the
    baseline has to degrade with it rather than vanish."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")
    sub = Submission(
        id=uuid.uuid4(),
        title="t",
        content_type="docx",
        file_path=str(path),
        original_content="Returns are guaranteed.",
        current_content="Returns are not guaranteed.",
    )
    assert lexical_document_service.import_text(sub) is None
    res = _diff(sub)
    assert res["edited"] is True and res["changed"] == 1


# --- html_to_text: the words, in order --------------------------------------


def test_html_to_text_keeps_words_apart_and_sentences_together():
    text = html_to_text(
        "<p>Insurance is the <strong>subject</strong> matter.</p>"
        "<ul><li>parent<ul><li>child</li></ul></li><li>second</li></ul>"
        "<table><tr><td><p>a</p></td><td><p>b</p></td></tr></table>"
    )
    assert "Insurance is the subject matter." in text
    # A block boundary is a boundary; nothing may be glued across one, and no
    # block's own wording may be swallowed by a nested one.
    assert text.split("\n\n") == [
        "Insurance is the subject matter.",
        "parent",
        "child",
        "second",
        "a",
        "b",
    ]
