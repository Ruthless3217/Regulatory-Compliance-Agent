"""DOCX -> HTML -> DOCX must preserve structure end to end.

lexical_import and lexical_export are tested separately against hand-written
HTML, and they agreed on an idealised shape that mammoth does not actually
emit: real cell markup is `<td><p>a</p></td>`, not `<td>a</td>`. A test that
only ever sees hand-written HTML cannot catch the two halves drifting apart,
so this one feeds the importer's genuine output straight into the exporter.
"""
import io

from docx import Document

from app.services.lexical_export import lexical_html_to_docx
from app.services.lexical_import import docx_to_html


def _brochure() -> bytes:
    doc = Document()
    doc.add_heading("Charges", level=2)
    doc.add_paragraph("Returns are not guaranteed.")
    doc.add_paragraph("first", style="List Bullet")
    doc.add_paragraph("second", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Premium"
    table.rows[0].cells[1].text = "12,000"
    table.rows[1].cells[0].text = "Term"
    table.rows[1].cells[1].text = "10 yrs"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _round_trip() -> Document:
    html = docx_to_html(_brochure())
    return Document(io.BytesIO(lexical_html_to_docx(html, "Brochure")))


def test_heading_survives_the_round_trip():
    out = _round_trip()
    styled = {(p.text, p.style.name) for p in out.paragraphs if p.text.strip()}
    assert ("Charges", "Heading 2") in styled


def test_body_paragraph_survives():
    out = _round_trip()
    assert any(p.text == "Returns are not guaranteed." for p in out.paragraphs)


def test_bullets_survive_as_list_paragraphs():
    out = _round_trip()
    bullets = [p.text for p in out.paragraphs if p.style.name == "List Bullet"]
    assert bullets == ["first", "second"]


def test_table_survives_mammoths_nested_cell_markup():
    """The specific divergence: mammoth wraps cell text in <p>."""
    out = _round_trip()
    assert len(out.tables) == 1
    rows = [[c.text for c in r.cells] for r in out.tables[0].rows]
    assert rows == [["Premium", "12,000"], ["Term", "10 yrs"]]


def test_importer_really_does_nest_cell_content():
    """Pins the assumption the exporter has to tolerate, so a future mammoth
    upgrade that changes it fails here rather than silently emptying tables."""
    html = docx_to_html(_brochure())
    assert "<td><p>" in html


# ---------------------------------------------------------------------------
# ...and with the upload as the template, the round trip must preserve the
# document's LOOK, not only its structure.
#
# The formatting is not in the HTML and never was — mammoth converts a DOCX to
# semantic HTML, so "Georgia 22pt #003399, centred, indented" is gone by the
# time the editor sees it. It survives only in the upload's own XML, which is
# why the export edits that XML instead of rebuilding from the HTML.
# ---------------------------------------------------------------------------

import os  # noqa: E402

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402


def _branded_upload(tmp_path) -> str:
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run("BAJAJ ALLIANZ LIFE GOAL ASSURE")
    run.font.name, run.font.size = "Georgia", Pt(22)
    run.font.color.rgb = RGBColor(0x00, 0x33, 0x99)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body = doc.add_paragraph()
    body_run = body.add_run("Guaranteed returns of 8% a year, every year.")
    body_run.font.name, body_run.font.size = "Calibri", Pt(11)
    body.paragraph_format.left_indent = Inches(0.5)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Premium"
    table.rows[0].cells[1].text = "12,000"

    path = str(tmp_path / "brochure.docx")
    doc.save(path)
    return path


def _corrected(path: str, before: str, after: str) -> Document:
    with open(path, "rb") as f:
        html = docx_to_html(f.read())
    assert before in html, "the fixture no longer says what the test corrects"
    return Document(
        io.BytesIO(lexical_html_to_docx(html.replace(before, after), "Brochure", path))
    )


def test_the_real_round_trip_keeps_the_documents_look(tmp_path):
    """One correction, everything else as it was — which is the whole job."""
    path = _branded_upload(tmp_path)
    out = _corrected(path, "returns of 8% a year", "returns of 6% a year")

    corrected = next(r for p in out.paragraphs for r in p.runs if "6% a year" in r.text)
    assert corrected.font.name == "Calibri"
    assert corrected.font.size == Pt(11)

    title = next(p for p in out.paragraphs if "GOAL ASSURE" in p.text)
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.name == "Georgia"
    assert title.runs[0].font.color.rgb == RGBColor(0x00, 0x33, 0x99)

    assert out.tables[0].style.name == "Light Grid Accent 1"


def test_the_upload_is_still_the_immutable_original(tmp_path):
    path = _branded_upload(tmp_path)
    before = os.path.getmtime(path)
    with open(path, "rb") as f:
        original = f.read()
    _corrected(path, "returns of 8% a year", "returns of 6% a year")
    with open(path, "rb") as f:
        assert f.read() == original
    assert os.path.getmtime(path) == before


def test_an_untouched_document_comes_back_formatted_character_for_character(tmp_path):
    """The strongest form of the promise, and the one that catches drift.

    Every visible character must still be set exactly as it was — same run
    properties, same values, written the same way. A bare ``<w:b/>`` rewritten
    as ``<w:b w:val="1"/>`` renders identically and is still a regression: it
    means the exporter is touching formatting it was not asked to change, and
    whatever does that to bold will eventually do it to something that shows.
    """
    import collections
    import re

    from docx.oxml.ns import qn

    def char_formats(document):
        counted = collections.Counter()
        for run in document.element.body.iter(qn("w:r")):
            rPr = run.find(qn("w:rPr"))
            key = re.sub(r"\s+", " ", rPr.xml.split(">", 1)[1]) if rPr is not None else ""
            for node in run.findall(qn("w:t")):
                for char in (node.text or ""):
                    if not char.isspace():
                        counted[(char, key)] += 1
        return counted

    path = _branded_upload(tmp_path)
    with open(path, "rb") as f:
        data = f.read()
    out = Document(io.BytesIO(lexical_html_to_docx(docx_to_html(data), "Brochure", path)))

    before = char_formats(Document(io.BytesIO(data)))
    assert before, "the fixture must have formatted text to compare"
    assert not (before - char_formats(out))
