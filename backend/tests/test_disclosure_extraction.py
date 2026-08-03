"""Submission-path extraction completeness (Past Performance root-cause fix).

The failing verdict happened because the creative's mandated disclaimer lived
outside the extracted surface: DOCX footers/headers/tables/text-boxes are not in
``Document.paragraphs``, and image-only PDF pages have no text layer in the
preprocessing path (OCR existed only in the Compare tool). These tests pin the
extraction fix: everything a consumer can SEE must reach the compliance text.
"""
import asyncio
import io
import sys
import types

import img2pdf
import pytest
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import parse_xml

from app.services.preprocessing_service import ContextEngineeringService

DISCLAIMER = "Past performance is not indicative of future performance."


def _svc():
    return ContextEngineeringService(db=None)


def _extract_docx(path):
    return asyncio.run(_svc()._extract_docx(str(path)))


# --- DOCX: non-body surfaces --------------------------------------------------

def test_docx_footer_disclaimer_is_extracted(tmp_path):
    doc = Document()
    doc.add_paragraph("Between 2022 and 2024 our equity-linked fund delivered 22% p.a.")
    doc.sections[0].footer.paragraphs[0].text = DISCLAIMER
    p = tmp_path / "footer.docx"
    doc.save(p)

    text = _extract_docx(p)
    assert DISCLAIMER in text
    assert "22% p.a." in text  # body still present


def test_docx_header_is_extracted(tmp_path):
    doc = Document()
    doc.add_paragraph("Body copy.")
    doc.sections[0].header.paragraphs[0].text = "Bajaj Life Insurance Limited"
    p = tmp_path / "header.docx"
    doc.save(p)

    assert "Bajaj Life Insurance Limited" in _extract_docx(p)


def test_docx_table_text_is_extracted(tmp_path):
    doc = Document()
    doc.add_paragraph("Fund options:")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Equity Growth Fund"
    table.cell(0, 1).text = DISCLAIMER
    p = tmp_path / "table.docx"
    doc.save(p)

    text = _extract_docx(p)
    assert "Equity Growth Fund" in text
    assert DISCLAIMER in text


_VML_TEXTBOX = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
    '<v:shape><v:textbox><w:txbxContent><w:p><w:r>'
    "<w:t>{text}</w:t>"
    "</w:r></w:p></w:txbxContent></v:textbox></v:shape>"
    "</w:pict></w:r></w:p>"
)


def test_docx_textbox_disclaimer_is_extracted(tmp_path):
    doc = Document()
    doc.add_paragraph("Invest in our Equity Growth Fund for long-term wealth.")
    doc.element.body.append(parse_xml(_VML_TEXTBOX.format(text=DISCLAIMER)))
    p = tmp_path / "textbox.docx"
    doc.save(p)

    assert DISCLAIMER in _extract_docx(p)


def test_docx_footer_not_duplicated_per_section(tmp_path):
    doc = Document()
    doc.add_paragraph("Body.")
    doc.sections[0].footer.paragraphs[0].text = DISCLAIMER
    p = tmp_path / "one_footer.docx"
    doc.save(p)

    text = _extract_docx(p)
    assert text.count(DISCLAIMER) == 1


# --- PDF: image-only pages get OCR --------------------------------------------

def _image_only_pdf(tmp_path, text="footer text"):
    img = Image.new("RGB", (1000, 300), (255, 255, 255))
    ImageDraw.Draw(img).text((40, 130), text, fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG", dpi=(150, 150))
    path = tmp_path / "scanned.pdf"
    path.write_bytes(img2pdf.convert(buf.getvalue()))
    return str(path)


def _stub_pytesseract(monkeypatch, returns):
    mod = types.ModuleType("pytesseract")
    mod.image_to_string = lambda pil, *a, **k: returns
    monkeypatch.setitem(sys.modules, "pytesseract", mod)


def test_pdf_image_only_uses_ocr_fallback(tmp_path, monkeypatch):
    pdf = _image_only_pdf(tmp_path)
    _stub_pytesseract(monkeypatch, DISCLAIMER)
    text = asyncio.run(_svc()._extract_pdf(pdf))
    assert DISCLAIMER in text


def test_pdf_image_only_without_engine_returns_empty_not_crash(tmp_path, monkeypatch):
    pdf = _image_only_pdf(tmp_path)
    monkeypatch.setitem(sys.modules, "pytesseract", None)
    text = asyncio.run(_svc()._extract_pdf(pdf))
    assert text == ""
