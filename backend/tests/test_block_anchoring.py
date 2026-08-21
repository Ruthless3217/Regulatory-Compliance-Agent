"""Stable violation anchoring: chunk over the editor's blocks, then name one.

Three things have to line up or a finding cannot be highlighted where it was
written:

  * the id the backend mints for a block must be the id the editor mints for
    the same block — the expected values below were produced by node's own
    crypto over the frontend's normalize(), which is exactly what
    `frontend/components/editor/__tests__/sectionMap.test.ts` pins sectionMap.ts against;
  * a chunk must be made of whole blocks, and must say which;
  * a finding may name a block only when that block really holds its quote —
    a wrong id sends the editor's tier-0 locator to the wrong paragraph, which
    is worse than no id at all.
"""
import uuid

import pytest

from app.services.lexical_anchor import block_id, block_ids, normalize
from app.services.lexical_import import html_to_blocks, html_to_text
from app.services.preprocessing_service import ContextEngineeringService


# ---------------------------------------------------------------- block ids ---
#
# Frontend-derived: node -e "createHash('sha1').update(text.replace(/\s+/g,' ')
# .trim().toLowerCase()).digest('hex').slice(0,12)".

FRONTEND_IDS = {
    "Charges": "c758ce5194d7",
    "Guaranteed 15% returns.": "679d013694b2",
    "  Terms   and\nconditions apply.  ": "7979826dd58d",
    "Bajaj Allianz Life Goal Assure": "5d1b9d285dee",
}


@pytest.mark.parametrize("text,expected", sorted(FRONTEND_IDS.items()))
def test_block_id_matches_the_editors(text, expected):
    assert block_id(text) == expected


def test_normalize_folds_whitespace_and_case_like_the_editor():
    assert normalize("  Terms   and\nconditions apply.  ") == "terms and conditions apply."
    # Which is why a reflow or a case tweak does not move the id.
    assert block_id("Charges") == block_id("charges") == block_id("  Charges\n")


def test_identical_blocks_are_disambiguated_by_ordinal():
    ids = block_ids(["Terms apply.", "Other copy.", "Terms apply.", "Terms apply."])
    base = block_id("Terms apply.")
    assert ids[0] == base
    assert ids[2] == f"{base}~1"
    assert ids[3] == f"{base}~2"
    assert len(set(ids)) == 4


def test_block_ids_are_positional_not_global():
    """The ordinal counts repeats SEEN SO FAR, so inserting a block above a
    duplicate does not renumber the ones below it out of step with the editor,
    which assigns them in the same single pass."""
    assert block_ids(["a", "b"]) == [block_id("a"), block_id("b")]
    assert block_ids(["a", "a", "b"])[:2] == [block_id("a"), block_id("a", 1)]


# ------------------------------------------------------------- html blocks ---

def test_top_level_elements_are_the_blocks():
    blocks = html_to_blocks(
        "<h2>Charges</h2><p>The <strong>fund</strong> management charge applies.</p>"
    )
    assert [b["tag"] for b in blocks] == ["h2", "p"]
    assert blocks[0]["text"] == "Charges"
    # Inline markup joins with NO invented space, as Lexical's projection does.
    assert blocks[1]["text"] == "The fund management charge applies."


def test_a_list_is_one_block_and_so_is_a_table():
    """$generateNodesFromDOM maps a whole <ul> onto one ListNode and a whole
    <table> onto one TableNode, so each is one block the editor can be
    anchored to — not one per item or cell."""
    blocks = html_to_blocks(
        "<ul><li>First item</li><li>Second item</li></ul>"
        "<table><tr><td>Year</td><td>Charge</td></tr></table>"
    )
    assert [b["tag"] for b in blocks] == ["ul", "table"]
    assert blocks[0]["text"] == "First item Second item"
    assert blocks[1]["text"] == "Year Charge"


def test_blank_blocks_are_dropped_like_the_editor_drops_them():
    blocks = html_to_blocks("<p>Real copy.</p><p>   </p><p><img src='x'></p><p>More.</p>")
    assert [b["text"] for b in blocks] == ["Real copy.", "More."]


def test_html_to_text_is_unchanged_by_the_block_split():
    html = "<h2>Charges</h2><ul><li>One</li><li>Two</li></ul><p>Body copy.</p>"
    assert html_to_text(html) == "Charges\n\nOne\n\nTwo\n\nBody copy."


# -------------------------------------------------------- which blocks, whose ---

def _docx_submission(path):
    import io

    from docx import Document

    from app.models.submission import Submission

    doc = Document()
    doc.add_heading("Charges", level=2)
    doc.add_paragraph("The fund management charge applies.")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))


def test_import_blocks_reads_the_upload_the_editor_is_seeded_from(tmp_path):
    from app.services import lexical_document_service as svc

    blocks = svc.import_blocks(_docx_submission(tmp_path / "a.docx"))
    assert [b["tag"] for b in blocks] == ["h2", "p"]
    assert blocks[0]["text"] == "Charges"


def test_the_saved_working_document_wins_over_the_upload(tmp_path):
    """Once the reviewer has saved, `lexical_html` IS the editor's document —
    minting ids from the upload would name blocks the editor no longer has."""
    from app.services import lexical_document_service as svc

    sub = _docx_submission(tmp_path / "a.docx")
    sub.lexical_html = "<h2>Charges</h2><p>The charge was rewritten by the reviewer.</p>"

    assert [b["text"] for b in svc.document_blocks(sub)] == [
        "Charges", "The charge was rewritten by the reviewer.",
    ]


def test_pasted_and_plain_submissions_keep_the_flat_chunker(tmp_path):
    """Their graded text (surfaced meta tags, [PAGE FOOTER] labels) is not what
    the editor holds, so block-aligning would change WHAT is graded."""
    from app.models.submission import Submission
    from app.services import lexical_document_service as svc

    for content_type in ("text", "html", "markdown"):
        sub = Submission(id=uuid.uuid4(), title="t", content_type=content_type,
                         original_content="body", file_path=None)
        sub.lexical_html = "<p>Even with a saved document.</p>"
        assert svc.document_blocks(sub) is None


# ---------------------------------------------------- block-aligned chunking ---

BODY = (
    "The fund management charge is levied daily and is deducted from the fund "
    "value of the policy at the applicable rate declared by the company for "
    "each of the investment funds available under this contract."
)


def _blocks(*pairs):
    return [{"tag": tag, "text": text} for tag, text in pairs]


def _svc():
    return ContextEngineeringService(db=None)


def test_chunks_are_made_of_whole_blocks_and_name_them():
    blocks = _blocks(("h2", "Charges"), ("p", BODY), ("h2", "Exclusions"), ("p", BODY))
    chunks = _svc()._chunk_by_blocks(blocks, "docx")

    ids = block_ids([b["text"] for b in blocks])
    assert [c["metadata"]["section_title"] for c in chunks] == ["Charges", "Exclusions"]
    assert [c["metadata"]["block_ids"] for c in chunks] == [ids[:2], ids[2:]]
    # Every chunk is exactly its blocks, joined — nothing is cut mid-block, and
    # nothing is dropped between chunks.
    assert [c["text"].split("\n\n") for c in chunks] == [
        ["Charges", BODY], ["Exclusions", BODY],
    ]


def test_a_heading_only_starts_a_new_chunk_once_the_last_one_is_worth_grading():
    """MIN_SECTION_TOKENS, unchanged from the flat chunker: a run of bare
    headings must not become a chunk each."""
    blocks = _blocks(("h2", "Charges"), ("h2", "Exclusions"), ("h2", "Benefits"))
    chunks = _svc()._chunk_by_blocks(blocks, "docx")
    assert len(chunks) == 1
    assert len(chunks[0]["metadata"]["block_ids"]) == 3


def test_a_preamble_before_the_first_heading_has_no_section_title():
    blocks = _blocks(("p", BODY), ("h2", "Charges"), ("p", BODY))
    chunks = _svc()._chunk_by_blocks(blocks, "docx")
    assert chunks[0]["metadata"]["section_title"] is None
    assert chunks[1]["metadata"]["section_title"] == "Charges"


def test_a_block_bigger_than_a_chunk_is_windowed_but_keeps_its_one_id():
    blocks = _blocks(("p", BODY * 60))
    chunks = _svc()._chunk_by_blocks(blocks, "docx")
    assert len(chunks) > 1
    only = block_ids([blocks[0]["text"]])[0]
    assert all(c["metadata"]["block_ids"] == [only] for c in chunks)


def test_block_chunking_is_deterministic():
    blocks = _blocks(("h2", "Charges"), ("p", BODY), ("p", BODY), ("h2", "Risks"),
                     ("ul", "One Two Three"), ("p", BODY))
    assert _svc()._chunk_by_blocks(blocks, "docx") == _svc()._chunk_by_blocks(blocks, "docx")


def test_heading_tags_are_headings_even_when_the_text_does_not_read_like_one():
    long_heading = "Charges applicable during the first five policy years, in detail"
    assert ContextEngineeringService._is_heading_line(long_heading) is False
    assert ContextEngineeringService._is_heading_block(
        {"tag": "h3", "text": long_heading}
    ) is True
    assert ContextEngineeringService._is_heading_block({"tag": "p", "text": BODY}) is False


# --------------------------------------------------------- anchor selection ---

from app.models.content_chunk import ContentChunk  # noqa: E402
from app.services.agents.compliance.engine import ComplianceEngine  # noqa: E402


def _chunk_row(texts, ids=None, title="Charges"):
    return ContentChunk(
        id=uuid.uuid4(),
        chunk_index=0,
        text="\n\n".join(texts),
        chunk_metadata={
            "section_title": title,
            "block_ids": ids if ids is not None else block_ids(texts),
        },
    )


def test_quote_found_in_a_block_gets_that_block_and_its_offsets():
    texts = ["Charges", "We offer guaranteed 15% returns on every policy."]
    row = _chunk_row(texts)

    out = ComplianceEngine._anchor_fields("Guaranteed 15% Returns", row)

    assert out["anchor_node_key"] == block_ids(texts)[1]
    assert out["section_title"] == "Charges"
    # Offsets index the block's NORMALIZED text — the frontend's coordinates.
    body = normalize(texts[1])
    assert body[out["anchor_offset_start"]:out["anchor_offset_end"]] == "guaranteed 15% returns"
    assert out["anchor_fingerprint"]


def test_the_fingerprint_is_the_surroundings_not_just_the_span():
    from app.services.lexical_anchor import compute_anchor_fingerprint

    texts = ["We offer guaranteed returns on every policy."]
    out = ComplianceEngine._anchor_fields("guaranteed returns", _chunk_row(texts))
    assert out["anchor_fingerprint"] == compute_anchor_fingerprint(
        "we offer ", "guaranteed returns", " on every policy."
    )


def test_a_quote_in_no_block_leaves_the_anchor_null():
    """NULL, never the chunk's first block: the editor scopes its search to the
    named block before verifying anything, so a wrong id turns a locatable
    finding into an unlocated one."""
    row = _chunk_row(["Charges", "The fund management charge applies."])
    out = ComplianceEngine._anchor_fields("guaranteed 15% returns", row)

    assert out["anchor_node_key"] is None
    assert out["anchor_offset_start"] is out["anchor_offset_end"] is None
    assert out["anchor_fingerprint"] is None
    # The section title still comes through — it is a property of the chunk.
    assert out["section_title"] == "Charges"


def test_a_chunk_with_no_block_ids_yields_only_a_section_title():
    row = ContentChunk(id=uuid.uuid4(), text="Guaranteed returns.",
                       chunk_metadata={"section_title": "Benefits"})
    out = ComplianceEngine._anchor_fields("guaranteed returns", row)
    assert out["section_title"] == "Benefits" and out["anchor_node_key"] is None


def test_a_document_level_finding_has_no_chunk_and_no_anchor():
    assert ComplianceEngine._anchor_fields("anything", None) == ComplianceEngine._NO_ANCHOR


def test_mismatched_metadata_says_nothing_rather_than_guessing():
    """Ids and text that do not line up describe a chunk this code did not
    build — the mapping from id to text would be a guess."""
    row = _chunk_row(["One block.", "Two blocks."], ids=["aaaaaaaaaaaa", "bbb", "ccc"])
    assert ComplianceEngine._anchor_fields("two blocks", row)["anchor_node_key"] is None

    # ...except the one legitimate mismatch: a single oversized block windowed
    # by tokens, where one id owns the whole slice.
    windowed = ContentChunk(
        id=uuid.uuid4(), text="Part one.\n\nPart two.",
        chunk_metadata={"block_ids": ["aaaaaaaaaaaa"], "section_title": None},
    )
    assert ComplianceEngine._anchor_fields("part two", windowed)["anchor_node_key"] == (
        "aaaaaaaaaaaa"
    )


def test_the_first_block_holding_the_quote_wins():
    texts = ["Terms apply.", "Terms apply."]
    row = _chunk_row(texts)
    assert ComplianceEngine._anchor_fields("terms apply", row)["anchor_node_key"] == (
        block_ids(texts)[0]
    )


# --------------------------------------------------- anchors through persist ---

class _Query:
    def __init__(self, rows):
        self._rows = rows

    def filter(self, *_a, **_kw):
        return self

    def all(self):
        return list(self._rows)

    def first(self):
        return self._rows[0] if self._rows else None

    def scalar(self):
        return None


class _PersistDB:
    """Enough Session for persist_results with no submission row (which is what
    keeps the render/anchor pass out of this test)."""

    def __init__(self, chunks=()):
        self._chunks = list(chunks)
        self.added = []
        self.commits = 0

    def query(self, model, *_rest):
        return _Query(self._chunks if model is ContentChunk else [])

    def add(self, obj):
        self.added.append(obj)

    def flush(self):
        pass

    def commit(self):
        self.commits += 1

    def refresh(self, _obj):
        pass

    def rollback(self):
        pass


def _persist(db, violations):
    ComplianceEngine.persist_results(
        submission_id=str(uuid.uuid4()), violations=violations, scores={}, db=db,
    )
    from app.models.violation import Violation

    return [o for o in db.added if isinstance(o, Violation)]


def test_persist_writes_the_anchor_onto_the_violation():
    texts = ["Charges", "We offer guaranteed 15% returns on every policy."]
    row = _chunk_row(texts)
    db = _PersistDB([row])

    written = _persist(db, [{
        "description": "guarantee", "current_text": "guaranteed 15% returns",
        "chunk_id": str(row.id), "chunk_index": 0,
    }])

    assert len(written) == 1
    assert written[0].anchor_node_key == block_ids(texts)[1]
    assert written[0].section_title == "Charges"
    assert written[0].anchor_offset_end > written[0].anchor_offset_start
    assert written[0].anchor_fingerprint


def test_persist_leaves_a_document_level_finding_unanchored():
    db = _PersistDB()
    written = _persist(db, [{"description": "missing ULIP disclaimer"}])
    assert written[0].anchor_node_key is None and written[0].section_title is None


def test_a_broken_anchor_never_costs_the_findings(monkeypatch):
    def _boom(*_a, **_kw):
        raise RuntimeError("chunk metadata is nonsense")

    monkeypatch.setattr(ComplianceEngine, "_anchor_fields", staticmethod(_boom))
    row = _chunk_row(["Guaranteed returns."])
    db = _PersistDB([row])

    written = _persist(db, [{"description": "guarantee", "current_text": "guaranteed returns",
                             "chunk_id": str(row.id)}])

    assert len(written) == 1 and db.commits == 1
    assert written[0].anchor_node_key is None


# ------------------------------------------------- reviewer rejections in the prompt ---

CONTENT = "Guaranteed 15% returns. Best plan in India."

PRECEDENT = {
    "id": "p1", "chunk_text": "Past copy about returns.",
    "anchor_text": "guaranteed returns", "comment_text": "Remove the guarantee.",
    "violation_category": "Misleading claim", "severity": "critical",
}
REJECTION = {
    "id": "p2", "chunk_text": "Past copy about tax.",
    "anchor_text": "tax free under Section 80C",
    "comment_text": "Standard wording, fine.",
    "violation_category": "Not a violation — Tax claim",
    "severity": "informational",
    "why_rationale": "A compliance reviewer read this exact finding and judged it "
                     "NOT a violation.",
}


def _prompt(precedents, **kw):
    return _svc().create_precedent_prompts(CONTENT, precedents, **kw)


def test_a_rejection_is_rendered_apart_from_the_precedents():
    out = _prompt([PRECEDENT, REJECTION])

    assert "REVIEWER REJECTIONS" in out
    assert "these were judged NOT violations" in out
    assert "--- REJECTION 1 ---" in out
    # ...and never as a precedent to cite.
    assert "--- PRECEDENT 1 ---" not in out
    assert "--- PRECEDENT 0 ---" in out
    assert "Phrase that was raised and REJECTED" in out


def test_the_rejection_instruction_forbids_flagging_and_citing():
    out = _prompt([REJECTION])
    assert "Do NOT flag" in out
    assert "do NOT cite them" in out


def test_precedent_indexes_survive_the_split():
    """`citations[].precedent_index` indexes the list as passed in, so a
    rejection ahead of a precedent must not renumber it."""
    out = _prompt([REJECTION, PRECEDENT])
    assert "--- PRECEDENT 1 ---" in out
    assert "--- REJECTION 0 ---" in out


def test_a_section_whose_only_precedents_are_rejections_has_nothing_to_cite():
    out = _prompt([REJECTION])
    assert "Do\nNOT emit any `citations`" in out
    assert "REVIEWER REJECTIONS" in out


def test_precedent_only_prompts_are_untouched_by_the_feature():
    out = _prompt([PRECEDENT])
    assert "REVIEWER REJECTIONS" not in out
    assert "--- PRECEDENT 0 ---" in out
    assert "Emit one `citations` entry per applicable" in out


def test_the_rejection_block_is_fenced_and_stable():
    first = _prompt([PRECEDENT, REJECTION])
    assert first == _prompt([PRECEDENT, REJECTION])
    # Fenced with the same content-derived delimiter as every untrusted block.
    fence = first.split("«")[1].split("»")[0]
    assert fence.startswith("UNTRUSTED-")
    header = first.index("REVIEWER REJECTIONS")
    assert first.index(f"«{fence}»", header) < first.index("--- REJECTION 1 ---")


def test_the_rejection_moves_the_fence():
    """The fence is derived from everything it fences; a rejection block that
    did not contribute could reuse the delimiter of a prompt without it."""
    assert _prompt([PRECEDENT, REJECTION]).split("«")[1] != _prompt([PRECEDENT]).split("«")[1]
