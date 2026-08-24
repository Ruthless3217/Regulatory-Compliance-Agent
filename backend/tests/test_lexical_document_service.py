"""Import must never mutate the upload, and must degrade rather than fail.

Degrade *visibly*: the bottom half of this file pins the three answers the
editor can get back, because a blank editable page looked identical for all of
them and the reviewer had no way to tell a scanned PDF from a broken one.
"""
import asyncio
import uuid

import pytest

from app.models.submission import Submission
from app.services import lexical_document_service as svc


def test_non_docx_returns_none_rather_than_raising(tmp_path):
    sub = Submission(id=uuid.uuid4(), title="t", content_type="text",
                     original_content="hello", file_path=None)
    assert svc.build_import_html(sub) is None


def test_missing_file_returns_none(tmp_path):
    sub = Submission(id=uuid.uuid4(), title="t", content_type="docx",
                     file_path=str(tmp_path / "absent.docx"))
    assert svc.build_import_html(sub) is None


def test_docx_returns_html_and_leaves_the_file_untouched(tmp_path):
    import io
    from docx import Document

    path = tmp_path / "a.docx"
    doc = Document()
    doc.add_heading("Charges", level=2)
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    before = path.read_bytes()

    sub = Submission(id=uuid.uuid4(), title="t", content_type="docx",
                     file_path=str(path))
    html = svc.build_import_html(sub)

    assert "<h2>" in html
    assert path.read_bytes() == before, "the upload is immutable"


def test_unconvertible_docx_returns_none(tmp_path):
    """A corrupt upload keeps the submission on extracted text rather than
    breaking GET /submissions/{id}."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")

    sub = Submission(id=uuid.uuid4(), title="t", content_type="docx",
                     file_path=str(path))
    assert svc.build_import_html(sub) is None


# --- can_import: answer the question without doing the work -----------------
#
# GET /submissions/{id} tells the editor whether there is something to seed
# from. Computing the HTML to find out cost seconds of CPU on a long document,
# inside an async handler, which both timed the request out and blocked the
# event loop for every other request on the worker.

def test_can_import_does_not_read_the_file(tmp_path, monkeypatch):
    import io as _io
    from docx import Document as _Document

    path = tmp_path / "big.docx"
    doc = _Document()
    doc.add_paragraph("body")
    buf = _io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())

    opened: list[str] = []
    real_open = open

    def spy(file, *a, **kw):
        opened.append(str(file))
        return real_open(file, *a, **kw)

    monkeypatch.setattr("builtins.open", spy)

    sub = Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))
    assert svc.can_import(sub) is True
    assert opened == [], "can_import must not open the upload"


def test_can_import_agrees_with_build_import_html(tmp_path):
    """A yes here that build_import_html then refuses would show the editor an
    empty document with no explanation."""
    import io as _io
    from docx import Document as _Document

    path = tmp_path / "a.docx"
    doc = _Document()
    doc.add_paragraph("body")
    buf = _io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())

    ok = Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))
    assert svc.can_import(ok) is True and svc.build_import_html(ok) is not None

    for bad in (
        Submission(id=uuid.uuid4(), title="t", content_type="text", file_path=str(path)),
        Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=None),
        Submission(id=uuid.uuid4(), title="t", content_type="docx",
                   file_path=str(tmp_path / "absent.docx")),
    ):
        assert svc.can_import(bad) is False
        assert svc.build_import_html(bad) is None


# --- import_html: which of the three answers -------------------------------
#
# `can_import` (and so `has_import_source`) only says the upload LOOKS
# importable — right content type, file on disk. A scanned PDF and a corrupt
# DOCX both pass it and fail here, and the reviewer used to be handed a blank
# editor with no explanation for either.


def _docx(path, *, empty=False):
    import io
    from docx import Document

    doc = Document()
    if not empty:
        doc.add_heading("Charges", level=2)
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))


def test_import_html_reports_success(tmp_path):
    result = svc.import_html(_docx(tmp_path / "a.docx"))
    assert result.status == "imported"
    assert "<h2>" in result.html
    assert result.reason is None


def test_import_html_reports_a_broken_conversion_with_a_reason(tmp_path):
    """`failed`, not `unavailable`: something went wrong, and the reviewer is
    owed the detail rather than a blank page."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")

    result = svc.import_html(
        Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))
    )
    assert result.status == "failed"
    assert result.html is None
    assert result.reason, "a failure with no reason is the bug being fixed"


def test_import_html_reports_nothing_to_import(tmp_path):
    """The three ways `can_import` says no, each with its own sentence."""
    for sub in (
        Submission(id=uuid.uuid4(), title="t", content_type="text", file_path=None),
        Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=None),
        Submission(id=uuid.uuid4(), title="t", content_type="pdf",
                   file_path=str(tmp_path / "absent.pdf")),
    ):
        result = svc.import_html(sub)
        assert result.status == "unavailable" and result.html is None
        assert result.reason


def test_empty_document_is_unavailable_not_a_failure(tmp_path):
    result = svc.import_html(_docx(tmp_path / "blank.docx", empty=True))
    assert result.status == "unavailable"
    assert result.html is None
    assert "no text" in result.reason


def test_scanned_pdf_is_a_property_of_the_document_not_a_failure(tmp_path):
    """A PDF with no text layer will never be editable here. Calling that a
    conversion failure sends the reviewer chasing a bug that does not exist —
    so it is `unavailable`, and the reason says scan/outlined artwork."""
    pytest.importorskip("reportlab")
    import io

    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.showPage()  # a page with no text at all
    c.save()
    path = tmp_path / "scan.pdf"
    path.write_bytes(buf.getvalue())

    result = svc.import_html(
        Submission(id=uuid.uuid4(), title="t", content_type="pdf", file_path=str(path))
    )
    assert result.status == "unavailable"
    assert result.html is None
    assert "text layer" in result.reason and "scan" in result.reason


def test_build_import_html_stays_best_effort(tmp_path):
    """The Optional[str] view is unchanged for callers that only degrade."""
    assert "<h2>" in svc.build_import_html(_docx(tmp_path / "a.docx"))

    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")
    assert svc.build_import_html(
        Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(path))
    ) is None


# --- GET /submissions/{id}/import-html --------------------------------------


class _FakeDB:
    """Enough Session for the route's single lookup."""

    def __init__(self, row):
        self._row = row

    def query(self, _model):
        return self

    def filter(self, *_exprs):
        return self

    def first(self):
        return self._row


class _Admin:
    """Admin short-circuits the visibility guard on its role alone, so the
    minimal _FakeDB needs no assignment table. These tests are about HTML
    import, not about who may see the document."""
    id = None
    role = "admin"


def _import_html_response(sub):
    from app.api.routes import submissions as routes

    return asyncio.run(
        routes.get_submission_import_html(str(sub.id), user=_Admin(), db=_FakeDB(sub))
    )


def test_endpoint_distinguishes_the_three_outcomes(tmp_path):
    ok = _import_html_response(_docx(tmp_path / "a.docx"))
    assert ok["status"] == "imported" and ok["html"]

    broken_path = tmp_path / "broken.docx"
    broken_path.write_bytes(b"not a docx at all")
    failed = _import_html_response(
        Submission(id=uuid.uuid4(), title="t", content_type="docx", file_path=str(broken_path))
    )
    assert failed["status"] == "failed" and failed["html"] is None and failed["reason"]

    nothing = _import_html_response(
        Submission(id=uuid.uuid4(), title="t", content_type="text", file_path=None)
    )
    assert nothing["status"] == "unavailable" and nothing["html"] is None and nothing["reason"]


def test_endpoint_never_reseeds_over_a_saved_working_document(tmp_path):
    sub = _docx(tmp_path / "a.docx")
    sub.lexical_state = {"root": {}}
    res = _import_html_response(sub)
    assert res["html"] is None
    assert res["status"] == "unavailable" and "already has" in res["reason"]
