"""Export is generated from the working document, not from the upload.

The uploaded file is the immutable original; the approved artifact is this.
Structure must survive the round trip or the editor's fidelity is pointless.

The upload is also the *template*: "put a docx in, get the same docx back"
means the export clones the original's styles, page setup, headers and footers
rather than building a default-styled Word file from scratch.
"""
import base64
import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from app.services.lexical_export import lexical_html_to_docx
from app.services.lexical_import import docx_to_html

# 1x1 transparent PNG — the smallest thing python-docx will accept as a picture.
_PNG_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def _paragraphs(data: bytes):
    return [p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]


def _styles(data: bytes):
    return [p.style.name for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]


def test_headings_export_as_word_headings():
    data = lexical_html_to_docx("<h2>Charges</h2><p>Body copy.</p>", "t")
    assert "Charges" in _paragraphs(data)
    assert any(s.startswith("Heading") for s in _styles(data))


def test_list_items_export_as_list_paragraphs():
    data = lexical_html_to_docx("<ul><li>first</li><li>second</li></ul>", "t")
    assert "first" in _paragraphs(data) and "second" in _paragraphs(data)


def test_table_exports_as_a_table():
    data = lexical_html_to_docx("<table><tr><td>a</td><td>b</td></tr></table>", "t")
    doc = Document(io.BytesIO(data))
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "a"


def test_empty_html_produces_a_valid_document():
    """A blank document must not raise — it exports as an empty file."""
    assert lexical_html_to_docx("", "t")


# ---------------------------------------------------------------------------
# Inline fidelity — the character-level formatting the editor shows
# ---------------------------------------------------------------------------

def _first(data: bytes, index: int = 0):
    return [p for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()][index]


def test_inline_markup_does_not_glue_words_together():
    """The reviewer's sentence must survive intact.

    An inline tag splits a paragraph into several strings; joining them without
    their spacing turns "Hello world now" into "Helloworldnow" — the document
    reads as corrupted even though every character is present.
    """
    data = lexical_html_to_docx("<p>Hello <strong>world</strong> now.</p>", "t")
    assert _first(data).text == "Hello world now."


def test_bold_italic_and_underline_survive_as_runs():
    html = "<p>plain <b>bold</b> <i>italic</i> <u>under</u></p>"
    runs = {r.text: r for r in _first(lexical_html_to_docx(html, "t")).runs if r.text.strip()}
    assert runs["bold"].bold is True
    assert runs["italic"].italic is True
    assert runs["under"].underline is True
    assert runs["plain "].bold is not True


def test_nested_emphasis_keeps_both_formats():
    html = "<p><strong>bold and <em>also italic</em></strong></p>"
    run = next(r for r in _first(lexical_html_to_docx(html, "t")).runs if "italic" in r.text)
    assert run.bold is True and run.italic is True


def test_line_break_stays_a_break():
    data = lexical_html_to_docx("<p>one<br/>two</p>", "t")
    assert _first(data).text == "one\ntwo"


def test_paragraph_alignment_survives():
    data = lexical_html_to_docx('<p style="text-align: center">Centred.</p>', "t")
    assert _first(data).alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_nested_list_keeps_its_level():
    html = "<ul><li>outer<ul><li>inner</li></ul></li></ul>"
    doc = Document(io.BytesIO(lexical_html_to_docx(html, "t")))
    styles = {p.text: p.style.name for p in doc.paragraphs if p.text.strip()}
    assert styles["outer"] == "List Bullet"
    assert styles["inner"] == "List Bullet 2"


def test_heading_keeps_its_inline_formatting():
    data = lexical_html_to_docx("<h2>Key <em>charges</em></h2>", "t")
    para = _first(data)
    assert para.text == "Key charges"
    assert next(r for r in para.runs if "charges" in r.text).italic is True


def test_table_cell_keeps_inline_formatting_and_spacing():
    html = "<table><tr><td>fund <b>value</b></td></tr></table>"
    cell = Document(io.BytesIO(lexical_html_to_docx(html, "t"))).tables[0].rows[0].cells[0]
    assert cell.text == "fund value"
    assert next(r for r in cell.paragraphs[0].runs if r.text == "value").bold is True


# ---------------------------------------------------------------------------
# Template cloning — the upload's own styling is the export's styling
# ---------------------------------------------------------------------------

def _template(tmp_path) -> str:
    """A DOCX with everything a default `Document()` would throw away."""
    doc = Document()

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Garamond"
    heading2.font.size = Pt(20)
    doc.styles.add_style("Disclaimer", WD_STYLE_TYPE.PARAGRAPH)

    section = doc.sections[0]
    # Whole/half inches only: Word stores these in twips, so 5.83" would not
    # compare equal to itself after a save/load round trip.
    section.page_width, section.page_height = Inches(5.5), Inches(8.5)
    section.left_margin = section.right_margin = Inches(1.75)

    section.header.paragraphs[0].text = "Bajaj Allianz Life"
    section.footer.paragraphs[0].text = "Insurance is the subject matter of solicitation."

    doc.add_heading("OLD TEMPLATE HEADING", level=2)
    doc.add_paragraph("OLD TEMPLATE BODY")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "OLD TEMPLATE TABLE"

    path = tmp_path / "upload.docx"
    doc.save(str(path))
    return str(path)


def _exported(tmp_path, html: str) -> Document:
    return Document(io.BytesIO(lexical_html_to_docx(html, "t", _template(tmp_path))))


def test_template_custom_styles_survive_the_round_trip(tmp_path):
    """The point of the whole feature: the export looks like the upload."""
    out = _exported(tmp_path, "<h2>Charges</h2><p>Body copy.</p>")
    assert out.styles["Heading 2"].font.name == "Garamond"
    assert out.styles["Heading 2"].font.size == Pt(20)
    assert "Disclaimer" in [s.name for s in out.styles]
    assert ("Charges", "Heading 2") in {(p.text, p.style.name) for p in out.paragraphs if p.text}


def test_template_body_content_is_replaced_not_appended(tmp_path):
    out = _exported(tmp_path, "<p>New approved copy.</p>")
    body = "\n".join(p.text for p in out.paragraphs)
    assert "New approved copy." in body
    assert "OLD TEMPLATE BODY" not in body
    assert "OLD TEMPLATE HEADING" not in body
    assert out.tables == []


def test_template_page_setup_survives_because_sectpr_is_kept(tmp_path):
    out = _exported(tmp_path, "<p>New approved copy.</p>")
    section = out.sections[0]
    assert (section.page_width, section.page_height) == (Inches(5.5), Inches(8.5))
    assert section.left_margin == Inches(1.75)
    # sectPr must be the *last* child of the body, or Word repairs the file.
    assert out.element.body[-1].tag.endswith("}sectPr")


def test_labelled_footer_goes_back_to_the_footer_not_the_body(tmp_path):
    """lexical_import appends the footer to the body so it can be reviewed and
    corrected. Writing it back as body text would relocate a mandated
    disclaimer out of the footer of the approved document."""
    html = (
        "<p>Body copy.</p>"
        "<h3>Page footer</h3><p>Insurance is the subject matter of solicitation. v2</p>"
    )
    out = _exported(tmp_path, html)
    assert "v2" in "\n".join(p.text for p in out.sections[0].footer.paragraphs)
    assert "v2" not in "\n".join(p.text for p in out.paragraphs)


def test_footer_text_keeps_its_spacing(tmp_path):
    """The footer carries the mandated disclaimer — it goes back word for word."""
    html = "<p>Body.</p><h3>Page footer</h3><p>Insurance is the <b>subject</b> matter.</p>"
    out = _exported(tmp_path, html)
    assert "Insurance is the subject matter." in "\n".join(
        p.text for p in out.sections[0].footer.paragraphs
    )


def test_labelled_header_goes_back_to_the_header(tmp_path):
    html = "<p>Body copy.</p><h3>Page header</h3><p>Bajaj Allianz Life Insurance</p>"
    out = _exported(tmp_path, html)
    assert "Bajaj Allianz Life Insurance" in "\n".join(
        p.text for p in out.sections[0].header.paragraphs
    )
    assert "Bajaj Allianz Life Insurance" not in "\n".join(p.text for p in out.paragraphs)


def test_unmappable_label_leaves_the_original_footer_alone(tmp_path):
    """A footer labelled for a section that does not exist is dropped, never
    duplicated into the body."""
    html = "<p>Body copy.</p><h3>Page footer 9</h3><p>Orphaned disclaimer.</p>"
    out = _exported(tmp_path, html)
    assert "Orphaned disclaimer." not in "\n".join(p.text for p in out.paragraphs)
    footer = "\n".join(p.text for p in out.sections[0].footer.paragraphs)
    assert footer == "Insurance is the subject matter of solicitation."


def test_base64_image_is_embedded(tmp_path):
    html = f'<p><img src="data:image/png;base64,{_PNG_B64}" /></p>'
    out = _exported(tmp_path, html)
    assert len(out.inline_shapes) == 1


def test_unreadable_image_is_skipped_not_fatal(tmp_path):
    html = '<p><img src="https://example.com/logo.png" />Body copy.</p><p><img /></p>'
    out = _exported(tmp_path, html)
    assert "Body copy." in "\n".join(p.text for p in out.paragraphs)
    assert len(out.inline_shapes) == 0


def test_missing_template_falls_back_to_a_default_document(tmp_path):
    data = lexical_html_to_docx("<h2>Charges</h2>", "t", str(tmp_path / "gone.docx"))
    assert "Charges" in _paragraphs(data)


def test_no_template_still_produces_a_valid_document():
    data = lexical_html_to_docx("<h2>Charges</h2><p>Body copy.</p>", "t", None)
    assert data[:2] == b"PK"
    assert ("Charges", "Heading 2") in {
        (p.text, p.style.name) for p in Document(io.BytesIO(data)).paragraphs if p.text
    }


def test_template_file_is_byte_identical_after_export(tmp_path):
    """The upload is immutable. It is opened read-only and cloned in memory."""
    path = _template(tmp_path)
    with open(path, "rb") as f:
        before = f.read()
    lexical_html_to_docx("<h2>New</h2><h3>Page footer</h3><p>New footer.</p>", "t", path)
    with open(path, "rb") as f:
        assert f.read() == before


# ---------------------------------------------------------------------------
# In-place body editing — the ORIGINAL body's own formatting is the export's
#
# Cloning the template carried styles.xml, sectPr and the headers across, but
# the body was still discarded and rebuilt from semantic HTML: every font,
# size, colour, indent and table style the upload set DIRECTLY on its body was
# lost, even when the reviewer changed nothing. These pin the other half of
# "put a docx in, get the same docx back" — the body is EDITED, not rebuilt.
# ---------------------------------------------------------------------------

from docx.shared import RGBColor  # noqa: E402


def _branded(tmp_path) -> str:
    """An upload formatted the way a real creative is: directly, on the runs
    and paragraphs, in a style Word never named."""
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run("BAJAJ ALLIANZ LIFE GOAL ASSURE")
    run.font.name, run.font.size = "Georgia", Pt(22)
    run.font.color.rgb = RGBColor(0x00, 0x33, 0x99)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)

    body = doc.add_paragraph()
    body_run = body.add_run("A unit linked plan with guaranteed returns of 8% a year.")
    body_run.font.name, body_run.font.size = "Calibri", Pt(11)
    body.paragraph_format.left_indent = Inches(0.5)

    doc.add_paragraph()  # a blank line the layout depends on

    table = doc.add_table(rows=2, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Policy year"
    table.rows[0].cells[1].text = "Allocation charge"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "6%"

    path = tmp_path / "branded.docx"
    doc.save(str(path))
    return str(path)


def _round_trip(template: str, html: str) -> Document:
    return Document(io.BytesIO(lexical_html_to_docx(html, "t", template)))


def _run_of(doc: Document, needle: str):
    return next(
        r
        for p in doc.paragraphs
        for r in p.runs
        if needle in r.text
    )


_UNCHANGED = (
    "<p>BAJAJ ALLIANZ LIFE GOAL ASSURE</p>"
    "<p>A unit linked plan with guaranteed returns of 8% a year.</p>"
    "<table><tr><td>Policy year</td><td>Allocation charge</td></tr>"
    "<tr><td>1</td><td>6%</td></tr></table>"
)


def test_direct_run_formatting_survives_an_untouched_export(tmp_path):
    """The reviewer changed nothing, so the file must come back as it went in."""
    out = _round_trip(_branded(tmp_path), _UNCHANGED)
    title = _run_of(out, "GOAL ASSURE")
    assert title.font.name == "Georgia"
    assert title.font.size == Pt(22)
    assert title.font.color.rgb == RGBColor(0x00, 0x33, 0x99)


def test_paragraph_formatting_survives_an_untouched_export(tmp_path):
    out = _round_trip(_branded(tmp_path), _UNCHANGED)
    title = next(p for p in out.paragraphs if "GOAL ASSURE" in p.text)
    body = next(p for p in out.paragraphs if "unit linked plan" in p.text)
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.paragraph_format.space_after == Pt(18)
    assert body.paragraph_format.left_indent == Inches(0.5)


def test_table_keeps_its_style(tmp_path):
    out = _round_trip(_branded(tmp_path), _UNCHANGED)
    assert out.tables[0].style.name == "Light Grid Accent 1"


def test_blank_layout_paragraphs_are_not_deleted(tmp_path):
    """mammoth drops blank paragraphs, so the editor never held them. Reading
    their absence as a reviewer deletion would re-flow the page."""
    out = _round_trip(_branded(tmp_path), _UNCHANGED)
    assert any(not p.text.strip() for p in out.paragraphs)


def test_corrected_wording_keeps_the_formatting_around_it(tmp_path):
    """The point of patching rather than rebuilding: the run that held "8%"
    is the run that now holds "6%", so it is still Calibri 11."""
    html = _UNCHANGED.replace("returns of 8% a year", "returns of 6% a year")
    out = _round_trip(_branded(tmp_path), html)
    corrected = _run_of(out, "6% a year")
    assert corrected.font.name == "Calibri"
    assert corrected.font.size == Pt(11)


def test_a_corrected_table_cell_keeps_the_table(tmp_path):
    html = _UNCHANGED.replace("<td>6%</td>", "<td>4%</td>")
    out = _round_trip(_branded(tmp_path), html)
    assert out.tables[0].style.name == "Light Grid Accent 1"
    assert out.tables[0].rows[1].cells[1].text == "4%"


def test_a_new_paragraph_is_set_like_the_one_it_follows(tmp_path):
    html = _UNCHANGED.replace(
        "</p><table>",
        "</p><p>Tax benefits are subject to change in tax laws.</p><table>",
    )
    out = _round_trip(_branded(tmp_path), html)
    added = next(p for p in out.paragraphs if "Tax benefits" in p.text)
    assert added.paragraph_format.left_indent == Inches(0.5)
    assert added.runs[0].font.name == "Calibri"


def test_a_deleted_paragraph_is_removed(tmp_path):
    html = _UNCHANGED.replace(
        "<p>A unit linked plan with guaranteed returns of 8% a year.</p>", ""
    )
    out = _round_trip(_branded(tmp_path), html)
    assert "unit linked plan" not in "\n".join(p.text for p in out.paragraphs)
    assert "GOAL ASSURE" in "\n".join(p.text for p in out.paragraphs)


def test_bold_applied_in_the_editor_reaches_the_export(tmp_path):
    html = _UNCHANGED.replace(
        "guaranteed returns", "<strong>guaranteed</strong> returns"
    )
    out = _round_trip(_branded(tmp_path), html)
    assert _run_of(out, "guaranteed").bold is True
    # ...on top of the run's own formatting, not instead of it.
    assert _run_of(out, "guaranteed").font.name == "Calibri"


def test_corrected_footer_keeps_its_own_formatting(tmp_path):
    """A mandated disclaimer is set small by hand. Re-typing the line as plain
    text would hand it back in the body font."""
    path = _branded(tmp_path)
    doc = Document(path)
    footer_run = doc.sections[0].footer.paragraphs[0].add_run(
        "Insurance is the subject matter of solicitation."
    )
    footer_run.font.size = Pt(7)
    doc.save(path)

    html = (
        _UNCHANGED
        + "<h3>Page footer</h3>"
        + "<p>Insurance is the subject matter of the solicitation.</p>"
    )
    out = Document(io.BytesIO(lexical_html_to_docx(html, "t", path)))
    corrected = out.sections[0].footer.paragraphs[0]
    assert corrected.text == "Insurance is the subject matter of the solicitation."
    assert corrected.runs[0].font.size == Pt(7)


# ---------------------------------------------------------------------------
# Text boxes — a mandated disclaimer is very often set in one
# ---------------------------------------------------------------------------

_TXBX = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    "<w:drawing><wp:inline><a:graphic><a:graphicData><wps:wsp><wps:txbx>"
    "<w:txbxContent><w:p><w:r><w:rPr><w:sz w:val=\"14\"/></w:rPr>"
    "<w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    "</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>"
)


def _with_textbox(tmp_path, text: str) -> str:
    """An upload whose disclaimer sits in a DrawingML text box — the shape
    mammoth walks into and does not convert, so lexical_import appends it to
    the body as a labelled section."""
    from lxml import etree

    doc = Document()
    doc.add_paragraph("Body copy.")
    doc.add_paragraph()._p.append(etree.fromstring(_TXBX.format(text=text)))
    path = tmp_path / "textbox.docx"
    doc.save(str(path))
    return str(path)


def test_text_box_label_is_not_written_into_the_body(tmp_path):
    """The label is scaffolding lexical_import adds so the box can be read in
    the editor. Exporting it leaves the word "Text box" in the document."""
    path = _with_textbox(tmp_path, "Tax benefits are subject to change.")
    html = "<p>Body copy.</p><h3>Text box</h3><p>Tax benefits are subject to change.</p>"
    out = _round_trip(path, html)
    body = "\n".join(p.text for p in out.paragraphs)
    assert "Text box" not in body


def test_text_box_wording_is_not_relocated_into_the_body(tmp_path):
    """Writing it out as a body paragraph moves the disclaimer out of the box
    it was designed into — and leaves the box saying the old thing."""
    path = _with_textbox(tmp_path, "Tax benefits are subject to change.")
    html = "<p>Body copy.</p><h3>Text box</h3><p>Tax benefits are subject to change.</p>"
    out = _round_trip(path, html)
    assert "Tax benefits" not in "\n".join(p.text for p in out.paragraphs)
    assert "Tax benefits are subject to change." in out.element.body.xml


def test_a_corrected_text_box_is_corrected_in_place(tmp_path):
    path = _with_textbox(tmp_path, "Tax benefits are subject to change.")
    html = (
        "<p>Body copy.</p><h3>Text box</h3>"
        "<p>Tax benefits are subject to change in tax laws.</p>"
    )
    out = _round_trip(path, html)
    xml = out.element.body.xml
    assert "Tax benefits are subject to change in tax laws." in xml
    assert 'w:val="14"' in xml  # the box's own type size, kept
    assert "Tax benefits" not in "\n".join(p.text for p in out.paragraphs)


def test_an_unmappable_text_box_label_is_dropped_not_written_to_the_body(tmp_path):
    path = _with_textbox(tmp_path, "Tax benefits are subject to change.")
    html = "<p>Body copy.</p><h3>Text box 9</h3><p>Orphaned disclaimer.</p>"
    out = _round_trip(path, html)
    body = "\n".join(p.text for p in out.paragraphs)
    assert "Orphaned disclaimer." not in body
    assert "Text box" not in body


def test_vertically_merged_cells_do_not_shift_the_row(tmp_path):
    """Word keeps an empty ``w:tc`` in every row a merge continues through;
    mammoth writes the merge as one ``rowspan`` and omits them. Pairing cells
    by raw position therefore slides a row's data one cell to the right — into
    the continuation cell, which Word requires to stay empty — from the first
    merged cell to the end of the table."""
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Rs. 25 Lakhs"
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).text = "1,75,825"
    table.cell(1, 1).text = "1,86,025"
    table.cell(2, 0).text = "Age at entry"
    table.cell(2, 1).text = "50 years"
    path = str(tmp_path / "merged.docx")
    doc.save(path)

    with open(path, "rb") as f:
        html = docx_to_html(f.read())
    out = Document(io.BytesIO(lexical_html_to_docx(html, "t", path)))

    # Read at the XML level: `row.cells` resolves a merge back to its origin,
    # so it reports the right answer even when the continuation cell has been
    # written into — which is exactly the corruption to catch.
    from docx.oxml.ns import qn

    rows = [
        [" ".join(t.text or "" for t in tc.iter(qn("w:t"))).strip()
         for tc in tr.findall(qn("w:tc"))]
        for tr in out.tables[0]._tbl.findall(qn("w:tr"))
    ]
    assert rows == [
        ["Rs. 25 Lakhs", "1,75,825"],
        ["", "1,86,025"],  # the continuation cell stays empty, as Word requires
        ["Age at entry", "50 years"],
    ]


def test_repeated_paragraphs_keep_their_own_formatting_when_one_is_edited(tmp_path):
    """These documents repeat a line verbatim — "Terms and conditions apply."
    under three different tables, each set differently.

    Block alignment is LCS-based, and LCS is free to match ANY equal pair. With
    repeats it will happily match the first original to the first edit and the
    *second* original to the *third* edit, reading the one real edit as an
    insert plus a delete somewhere else — which drops a paragraph's formatting
    on the floor and shifts the rest up one.
    """
    doc = Document()
    for size, name in ((9, "Georgia"), (14, "Calibri"), (20, "Arial")):
        run = doc.add_paragraph().add_run("Terms and conditions apply.")
        run.font.size, run.font.name = Pt(size), name
    path = str(tmp_path / "repeated.docx")
    doc.save(path)

    out = _round_trip(path, (
        "<p>Terms and conditions apply.</p>"
        "<p>Terms and conditions may apply.</p>"  # only the MIDDLE one edited
        "<p>Terms and conditions apply.</p>"
    ))

    written = [
        (p.text, p.runs[0].font.size, p.runs[0].font.name)
        for p in out.paragraphs if p.text.strip()
    ]
    assert written == [
        ("Terms and conditions apply.", Pt(9), "Georgia"),
        ("Terms and conditions may apply.", Pt(14), "Calibri"),
        ("Terms and conditions apply.", Pt(20), "Arial"),
    ]


def test_an_unchanged_picture_does_not_leave_a_blank_paragraph_behind(tmp_path):
    """mammoth hands the document's OWN pictures back as base64 `data:` URIs,
    so an image-only `<p>` in the editor is usually just the picture that is
    already sitting in the body untouched.

    Treating it as a block to insert adds an empty paragraph on every export —
    and that paragraph is cloned from its neighbour, so in a numbered document
    it inherits `numPr` and Word renders it as an empty bullet.
    """
    doc = Document()
    doc.add_paragraph("Body copy.", style="List Number")
    doc.add_picture(io.BytesIO(base64.b64decode(_PNG_B64)))
    doc.add_paragraph("More body copy.", style="List Number")
    path = str(tmp_path / "pictured.docx")
    doc.save(path)

    with open(path, "rb") as f:
        html = docx_to_html(f.read())
    assert "<img" in html, "the fixture must round-trip a picture through mammoth"
    out = _round_trip(path, html)

    assert len(out.inline_shapes) == 1, "the picture must not be embedded twice"
    assert len(out.paragraphs) == len(Document(path).paragraphs)
