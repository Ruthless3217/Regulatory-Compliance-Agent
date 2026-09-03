"""submission_export_service + export_common — the 9 export kinds behind
GET /submissions/{id}/export/{kind}.

Docx builders are exercised directly against real (unsaved) model instances,
matching test_submission_revisions_and_comments.py's style. The *.pdf/
bundle.zip path is tested with `export_common.docx_bytes_to_pdf` monkeypatched
(no network) — see test_gotenberg_client.py for the mocked-HTTP-transport
tests of the underlying Gotenberg calls, and the live-Docker smoke check
(run manually, documented in the task report) for a real end-to-end pass
against a running `compliance-gotenberg` container.
"""
import uuid
from unittest.mock import MagicMock

import pytest
from docx import Document
import io

from app.models.compliance_check import ComplianceCheck
from app.models.rule_feedback import RuleFeedback
from app.models.submission import Submission
from app.models.violation import Violation
from app.services import export_common as ec
from app.services import submission_export_service as ses


def _violation(**overrides) -> Violation:
    defaults = dict(
        id=uuid.uuid4(),
        compliance_check_id=uuid.uuid4(),
        category="disclosure",
        severity="high",
        description="Guaranteed-returns claim with no disclaimer.",
        current_text="guaranteed returns",
        suggested_fix="Add the standard market-risk disclaimer.",
        suppressed=False,
    )
    defaults.update(overrides)
    return Violation(**defaults)


def _feedback(**overrides) -> RuleFeedback:
    defaults = dict(
        id=uuid.uuid4(),
        violation_id=uuid.uuid4(),
        verdict="correct",
        reason="wrong_severity",
        original_text="guaranteed returns",
        final_text="potential returns (market-linked)",
        comment="Downgraded after legal review.",
    )
    defaults.update(overrides)
    return RuleFeedback(**defaults)


def _submission(text: str, **overrides) -> Submission:
    defaults = dict(id=uuid.uuid4(), title="Brochure v1", content_type="text", original_content=text)
    defaults.update(overrides)
    return Submission(**defaults)


# ---------------------------------------------------------------------------
# export_common: normalize_severity / find_spans / iter_paragraphs
# ---------------------------------------------------------------------------

def test_normalize_severity_maps_aliases_and_unknowns():
    assert ec.normalize_severity("CRITICAL") == "critical"
    assert ec.normalize_severity("moderate") == "medium"
    assert ec.normalize_severity("informational") == "low"
    assert ec.normalize_severity("info") == "low"
    assert ec.normalize_severity("bogus") == "medium"
    assert ec.normalize_severity(None) == "medium"


def test_find_spans_matches_case_and_whitespace_insensitively():
    text = "Enjoy   GUARANTEED\nreturns on your investment today."
    v = _violation(current_text="guaranteed returns")
    spans = ec.find_spans(text, [v])
    assert len(spans) == 1
    s = spans[0]
    assert text[s.start:s.end] == "GUARANTEED\nreturns"
    assert s.severity == "high"


def test_find_spans_finds_every_occurrence():
    text = "Repeat repeat REPEAT this phrase, then repeat this phrase again."
    v = _violation(current_text="repeat this phrase")
    spans = ec.find_spans(text, [v])
    assert len(spans) == 2


def test_find_spans_overlap_resolution_prefers_higher_severity():
    text = "This plan offers guaranteed high returns forever."
    short = _violation(current_text="guaranteed high returns", severity="low")
    long_critical = _violation(current_text="guaranteed high returns forever", severity="critical")
    spans = ec.find_spans(text, [short, long_critical])
    assert len(spans) == 1
    assert spans[0].severity == "critical"
    assert spans[0].violation is long_critical


def test_find_spans_ignores_empty_current_text():
    v = _violation(current_text="   ")
    assert ec.find_spans("anything here", [v]) == []


def test_iter_paragraphs_splits_on_blank_lines():
    text = "Para one line one\nline two.\n\nPara two.\n\n\nPara three."
    bounds = ec.iter_paragraphs(text)
    paras = [text[s:e] for s, e in bounds]
    assert paras == ["Para one line one\nline two.", "Para two.", "Para three."]


def test_iter_paragraphs_single_paragraph_has_no_split():
    text = "Just one paragraph, no blank lines."
    assert ec.iter_paragraphs(text) == [(0, len(text))]


# ---------------------------------------------------------------------------
# DOCX builders (called directly on real, unsaved model instances)
# ---------------------------------------------------------------------------

def test_clean_docx_contains_the_document_text():
    sub = _submission("First paragraph.\n\nSecond paragraph.")
    data = ses._clean_docx(sub)
    assert data[:2] == b"PK"
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "First paragraph." in body
    assert "Second paragraph." in body


def test_clean_docx_prefers_current_content_over_original():
    sub = _submission("stale original", current_content="fresh edited text")
    doc = Document(io.BytesIO(ses._clean_docx(sub)))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "fresh edited text" in body
    assert "stale original" not in body


def test_clean_docx_empty_content_shows_notice_not_blank():
    sub = _submission("")
    doc = Document(io.BytesIO(ses._clean_docx(sub)))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "No document content available" in body


def test_clean_docx_is_generated_from_the_working_document_when_there_is_one():
    """The approved artifact comes from what the reviewer edited, not from the
    extracted text — structure the editor showed has to survive."""
    sub = _submission(
        "Charges\n\nBody copy.",
        lexical_html="<h2>Charges</h2><p>Body copy.</p>",
    )
    doc = Document(io.BytesIO(ses._clean_docx(sub)))
    styled = {(p.text, p.style.name) for p in doc.paragraphs if p.text.strip()}
    assert ("Charges", "Heading 2") in styled


def test_clean_docx_working_document_wins_over_stale_extracted_text():
    sub = _submission(
        "Guaranteed returns.",
        current_content="Guaranteed returns.",
        lexical_html="<p>Returns are not guaranteed.</p>",
    )
    body = "\n".join(p.text for p in Document(io.BytesIO(ses._clean_docx(sub))).paragraphs)
    assert "Returns are not guaranteed." in body
    assert "Guaranteed returns." not in body


def test_clean_docx_leaves_the_uploaded_file_untouched(tmp_path):
    """The upload is the immutable original. Export never opens it for writing."""
    path = tmp_path / "upload.docx"
    src = Document()
    src.add_paragraph("Guaranteed returns.")
    buf = io.BytesIO()
    src.save(buf)
    path.write_bytes(buf.getvalue())
    before = path.read_bytes()

    sub = _submission(
        "Guaranteed returns.",
        content_type="docx",
        file_path=str(path),
        current_content="Returns are not guaranteed.",
        lexical_html="<p>Returns are not guaranteed.</p>",
    )
    ses._clean_docx(sub)
    assert path.read_bytes() == before


def test_clean_docx_uses_the_upload_as_the_export_template(tmp_path):
    """"Put a docx in, get the same docx back" — the upload's own styling and
    page setup have to come back out with the corrected text."""
    from docx.shared import Inches, Pt

    path = tmp_path / "upload.docx"
    src = Document()
    src.styles["Heading 2"].font.name = "Garamond"
    src.styles["Heading 2"].font.size = Pt(20)
    src.sections[0].left_margin = Inches(1.75)
    src.sections[0].footer.paragraphs[0].text = "Insurance is the subject matter of solicitation."
    src.add_paragraph("Guaranteed returns.")
    src.save(str(path))

    sub = _submission(
        "Guaranteed returns.",
        content_type="docx",
        file_path=str(path),
        lexical_html="<h2>Charges</h2><p>Returns are not guaranteed.</p>",
    )
    out = Document(io.BytesIO(ses._clean_docx(sub)))
    assert out.styles["Heading 2"].font.name == "Garamond"
    assert out.sections[0].left_margin == Inches(1.75)
    assert "solicitation" in "\n".join(p.text for p in out.sections[0].footer.paragraphs)
    body = "\n".join(p.text for p in out.paragraphs)
    assert "Returns are not guaranteed." in body
    assert "Guaranteed returns." not in body


def test_clean_docx_ignores_a_missing_or_non_docx_upload(tmp_path):
    """A pasted-text submission has no template; a stale file_path is not one
    either. Both still export."""
    for overrides in (
        dict(content_type="html", file_path=str(tmp_path / "not.docx")),
        dict(content_type="docx", file_path=str(tmp_path / "gone.docx")),
        dict(content_type="docx", file_path=None),
    ):
        sub = _submission("x", lexical_html="<h2>Charges</h2>", **overrides)
        body = "\n".join(p.text for p in Document(io.BytesIO(ses._clean_docx(sub))).paragraphs)
        assert "Charges" in body, overrides


def test_clean_docx_without_a_working_document_still_rebuilds_from_text():
    """Every submission predating the editor has lexical_html NULL."""
    sub = _submission("First paragraph.", lexical_html=None)
    body = "\n".join(p.text for p in Document(io.BytesIO(ses._clean_docx(sub))).paragraphs)
    assert "First paragraph." in body


def test_annotated_docx_highlights_and_numbers_findings():
    text = "This plan offers guaranteed returns on your money."
    v = _violation(current_text="guaranteed returns", severity="critical")
    data = ses._annotated_docx(_submission(text), [v])
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "[1]" in body  # numbered marker next to the highlighted span
    assert any("Findings" in p.text for p in doc.paragraphs)
    table = doc.tables[0]
    header = [c.text for c in table.rows[0].cells]
    assert header == ["#", "Severity", "Category", "Description"]
    assert table.rows[1].cells[1].text == "critical"


def test_annotated_docx_with_no_violations_has_no_findings_table():
    data = ses._annotated_docx(_submission("Plain, unremarkable copy."), [])
    doc = Document(io.BytesIO(data))
    assert doc.tables == []


def test_report_docx_lists_violations_with_suppressed_flag():
    violations = [
        _violation(severity="high", suppressed=False),
        _violation(severity="low", suppressed=True, description="Heading, not a claim"),
    ]
    check = ComplianceCheck(id=uuid.uuid4(), overall_score=72.5, grade="C")
    data = ses._report_docx(_submission("irrelevant"), check, violations)
    doc = Document(io.BytesIO(data))
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == [
        "#", "Severity", "Category", "Flagged text", "Description", "Suggested fix", "Suppressed",
    ]
    assert len(table.rows) == 1 + 2
    assert table.rows[1].cells[6].text == "No"
    assert table.rows[2].cells[6].text == "Yes"
    assert any("Score 72.5" in p.text for p in doc.paragraphs)


def test_report_docx_with_no_check_says_not_yet_analyzed():
    data = ses._report_docx(_submission("irrelevant"), None, [])
    doc = Document(io.BytesIO(data))
    assert any("No analysis run yet" in p.text for p in doc.paragraphs)


def test_feedback_report_docx_lists_reviewer_actions():
    fb = [_feedback(verdict="not_violation", reason="out_of_scope")]
    data = ses._feedback_report_docx(_submission("irrelevant"), fb)
    doc = Document(io.BytesIO(data))
    table = doc.tables[0]
    assert table.rows[1].cells[1].text == "not_violation"
    assert table.rows[1].cells[2].text == "out_of_scope"


# ---------------------------------------------------------------------------
# build_export dispatch (DB access mocked; PDF conversion mocked)
# ---------------------------------------------------------------------------

def _db_for(check, violations, feedback):
    db = MagicMock()

    def query_side_effect(model):
        q = MagicMock()
        if model is ComplianceCheck:
            q.filter.return_value.order_by.return_value.first.return_value = check
        elif model is Violation:
            q.filter.return_value.order_by.return_value.all.return_value = violations
        elif model is RuleFeedback:
            q.filter.return_value.order_by.return_value.all.return_value = feedback
        return q

    db.query.side_effect = query_side_effect
    return db


def test_build_export_docx_kinds_need_no_gotenberg():
    sub = _submission("Some clean copy.")
    db = _db_for(check=None, violations=[], feedback=[])
    for kind in ("clean.docx", "annotated.docx", "report.docx", "feedback-report.docx"):
        data = ses.build_export(db, sub, kind)
        assert data[:2] == b"PK", kind


def test_build_export_pdf_kinds_go_through_the_docx_to_pdf_seam(monkeypatch):
    calls = []

    def fake_convert(docx_bytes):
        calls.append(docx_bytes)
        return b"%PDF-fake"

    monkeypatch.setattr(ec, "docx_bytes_to_pdf", fake_convert)
    sub = _submission("Some clean copy.")
    db = _db_for(check=None, violations=[], feedback=[])

    for kind in ("clean.pdf", "annotated.pdf", "report.pdf", "feedback-report.pdf"):
        assert ses.build_export(db, sub, kind) == b"%PDF-fake"
    assert len(calls) == 4
    assert all(c[:2] == b"PK" for c in calls)  # each was handed real docx bytes


def test_build_export_bundle_zip_contains_all_eight_files(monkeypatch):
    monkeypatch.setattr(ec, "docx_bytes_to_pdf", lambda b: b"%PDF-fake")
    sub = _submission("Some clean copy.")
    db = _db_for(check=None, violations=[], feedback=[])

    import zipfile
    data = ses.build_export(db, sub, "bundle.zip")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
    assert names == {
        "clean.docx", "clean.pdf",
        "annotated.docx", "annotated.pdf",
        "report.docx", "report.pdf",
        "feedback-report.docx", "feedback-report.pdf",
    }


def test_build_export_bundle_zip_survives_a_pdf_conversion_failure(monkeypatch):
    def boom(docx_bytes):
        raise RuntimeError("gotenberg unreachable")

    monkeypatch.setattr(ec, "docx_bytes_to_pdf", boom)
    sub = _submission("Some clean copy.")
    db = _db_for(check=None, violations=[], feedback=[])

    import zipfile
    data = ses.build_export(db, sub, "bundle.zip")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
    # docx set still ships; no .pdf member made it in.
    assert names == {"clean.docx", "annotated.docx", "report.docx", "feedback-report.docx"}


def test_build_export_unknown_kind_raises_value_error():
    sub = _submission("x")
    db = _db_for(check=None, violations=[], feedback=[])
    with pytest.raises(ValueError):
        ses.build_export(db, sub, "not-a-real-kind")


# ---------------------------------------------------------------------------
# Degraded / partial analysis state in the exported artifacts.
#
# The export is what LEAVES the system. A run graded on incomplete evidence
# (engine.evaluate_persistability -> completed_with_warnings) previously
# exported as "Score 100 · Grade A · 0 finding(s)" with nothing to say the
# grade covered less than the whole document, and the export gate consulted
# only `findings_are_stale`. A reader of that file could not tell a partial
# result from a clean one.
# ---------------------------------------------------------------------------

WARNINGS = [
    {"code": "rider_uins_without_fact_cards",
     "detail": {"uins": ["116N216V01"], "chunk_indexes": [24]},
     "explanation": "a rider named in this document has no authoritative record"},
    {"code": "precedent_evidence_unavailable", "detail": {"precedents": 0}},
]


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts.extend(c.text for row in table.rows for c in row.cells)
    return "\n".join(parts)


def test_report_docx_states_that_a_warned_grade_is_partial():
    check = ComplianceCheck(id=uuid.uuid4(), overall_score=100.0, grade="A")
    text = _docx_text(
        ses._report_docx(_submission("irrelevant"), check, [], analysis_warnings=WARNINGS)
    )

    assert "incomplete evidence" in text.lower()
    assert "116N216V01" in text


def test_report_docx_names_every_limitation_not_just_the_first():
    check = ComplianceCheck(id=uuid.uuid4(), overall_score=100.0, grade="A")
    text = _docx_text(
        ses._report_docx(_submission("irrelevant"), check, [], analysis_warnings=WARNINGS)
    )

    assert "rider_uins_without_fact_cards" in text
    assert "precedent_evidence_unavailable" in text


def test_report_docx_of_a_fully_grounded_run_says_nothing_extra():
    check = ComplianceCheck(id=uuid.uuid4(), overall_score=100.0, grade="A")
    text = _docx_text(ses._report_docx(_submission("irrelevant"), check, []))

    assert "incomplete evidence" not in text.lower()


def test_annotated_docx_carries_the_partial_state_too():
    """The findings list is as much a claim of completeness as the score is."""
    text = _docx_text(
        ses._annotated_docx(_submission("guaranteed returns"), [_violation()],
                            analysis_warnings=WARNINGS)
    )

    assert "incomplete evidence" in text.lower()


def test_build_export_passes_the_runs_warnings_into_the_report():
    """End to end through the dispatcher: the warnings live on the AnalysisRun
    that produced the check, so build_export has to go and fetch them."""
    from app.models.analysis_run import AnalysisRun

    check = ComplianceCheck(id=uuid.uuid4(), overall_score=100.0, grade="A")
    run = AnalysisRun(
        id=uuid.uuid4(), submission_id=uuid.uuid4(), run_number=1,
        status="completed", compliance_check_id=check.id,
        run_metadata={"analysis_warnings": WARNINGS},
    )
    db = MagicMock()

    def _query(model):
        q = MagicMock()
        q.filter.return_value = q
        q.order_by.return_value = q
        q.first.return_value = check if model is ComplianceCheck else run
        q.all.return_value = []
        return q

    db.query.side_effect = _query
    text = _docx_text(ses.build_export(db, _submission("irrelevant"), "report.docx"))

    assert "incomplete evidence" in text.lower()
