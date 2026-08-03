"""clean.docx must carry accepted edits without flattening the upload.

The rebuild path discards fonts/tables/images, so these pin the cases where the
original file is edited in place and the cases that must refuse and fall back
rather than write to the wrong paragraph.
"""
import io
import uuid

import pytest
from docx import Document

from app.models.submission import Submission
from app.services import submission_export_service as svc


def _docx(paragraphs) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _texts(data: bytes):
    return [p.text for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]


@pytest.fixture
def upload(tmp_path):
    def make(paragraphs, original=None, current=None):
        path = tmp_path / f"{uuid.uuid4()}.docx"
        path.write_bytes(_docx(paragraphs))
        joined = "\n\n".join(paragraphs)
        return Submission(
            id=uuid.uuid4(),
            title="Brochure",
            content_type="docx",
            file_path=str(path),
            original_content=original if original is not None else joined,
            current_content=current,
        )
    return make


def test_unedited_docx_is_returned_byte_for_byte(upload):
    sub = upload(["Guaranteed returns.", "Terms apply."])
    with open(sub.file_path, "rb") as f:
        assert svc._clean_docx(sub) == f.read()


def test_edited_paragraph_is_rewritten_in_the_original(upload):
    sub = upload(
        ["Guaranteed returns.", "Terms apply."],
        current="Returns are not guaranteed.\n\nTerms apply.",
    )
    assert _texts(svc._clean_docx(sub)) == ["Returns are not guaranteed.", "Terms apply."]


def test_untouched_paragraphs_keep_their_run_formatting(upload):
    """The point of editing in place: unedited content is never written to."""
    sub = upload(
        ["Bold claim here.", "Small print."],
        current="Softened claim here.\n\nSmall print.",
    )
    doc = Document(io.BytesIO(svc._clean_docx(sub)))
    untouched = [p for p in doc.paragraphs if p.text == "Small print."][0]
    assert untouched.runs, "paragraph should still have its original run"


def test_added_paragraph_falls_back_to_rebuild(upload):
    """No anchor in the original file for a paragraph that never existed."""
    sub = upload(
        ["Guaranteed returns."],
        current="Guaranteed returns.\n\nNewly added disclaimer.",
    )
    assert svc._edited_original_docx(sub) is None
    assert "Newly added disclaimer." in " ".join(_texts(svc._clean_docx(sub)))


def test_ambiguous_repeated_paragraph_falls_back(upload):
    """Same text twice: editing either one would be a guess."""
    sub = upload(
        ["Terms apply.", "Terms apply."],
        current="Terms and conditions apply.\n\nTerms apply.",
    )
    assert svc._edited_original_docx(sub) is None


def test_non_docx_upload_uses_rebuild(upload):
    sub = Submission(
        id=uuid.uuid4(), title="Pasted", content_type="text",
        original_content="Some copy.", current_content=None,
    )
    assert svc._edited_original_docx(sub) is None
    assert svc._clean_docx(sub)


def test_missing_file_falls_back_instead_of_raising(upload):
    sub = upload(["Guaranteed returns."], current="Returns vary.")
    import os
    os.remove(sub.file_path)
    assert svc._edited_original_docx(sub) is None
    assert "Returns vary." in " ".join(_texts(svc._clean_docx(sub)))
