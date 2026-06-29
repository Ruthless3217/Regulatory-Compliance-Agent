"""Generate a formatted DOCX of all compliance rules grouped by category.

Reads backend/exports/all_rules_by_category.csv and writes
backend/exports/all_rules_by_category.docx with one table per rule category,
each row showing rule text, source and severity (sorted critical -> low).
"""
import csv
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches

EXPORTS = Path(__file__).resolve().parent.parent / "exports"
CSV_PATH = EXPORTS / "all_rules_by_category.csv"
DOCX_PATH = EXPORTS / "all_rules_by_category.docx"

# Bajaj brand palette
BAJAJ_BLUE = RGBColor(0x00, 0x33, 0x91)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x66, 0x66, 0x66)

SEV_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}
SEV_COLOR = {
    "critical": "C0392B",
    "high": "E67E22",
    "medium": "2E86C1",
    "low": "7F8C8D",
}
CATEGORY_ORDER = ["IRDAI Rules", "SEBI Rules", "Compliance Rules", "Brand Rules"]


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, color=None, size=9, align=None,
                  italic=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.text = ""
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color
    return para


def set_col_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = w


def load_rules():
    by_cat = defaultdict(list)
    with open(CSV_PATH, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            by_cat[row["rule_category"]].append(row)
    for cat in by_cat:
        by_cat[cat].sort(
            key=lambda r: (SEV_ORDER.get(r["severity"], 9),
                           r["source"].lower())
        )
    return by_cat


def build():
    by_cat = load_rules()
    total = sum(len(v) for v in by_cat.values())

    doc = Document()
    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    section = doc.sections[0]
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Regulatory Compliance Rulebook")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BAJAJ_BLUE
    run.font.name = "Calibri"

    sub = doc.add_paragraph()
    srun = sub.add_run(
        f"All active rules grouped by category  -  {total} rules total"
    )
    srun.font.size = Pt(11)
    srun.font.color.rgb = GREY
    srun.italic = True

    # Summary line per category
    summ = doc.add_paragraph()
    parts = []
    for cat in CATEGORY_ORDER:
        if cat in by_cat:
            parts.append(f"{cat}: {len(by_cat[cat])}")
    srun = summ.add_run("   |   ".join(parts))
    srun.font.size = Pt(10)
    srun.font.color.rgb = DARK
    srun.bold = True

    ordered = [c for c in CATEGORY_ORDER if c in by_cat]
    ordered += [c for c in by_cat if c not in CATEGORY_ORDER]

    widths = [Inches(4.3), Inches(2.3), Inches(0.9)]

    for cat in ordered:
        rules = by_cat[cat]
        doc.add_paragraph()
        head = doc.add_paragraph()
        hrun = head.add_run(f"{cat}  ({len(rules)})")
        hrun.bold = True
        hrun.font.size = Pt(15)
        hrun.font.color.rgb = BAJAJ_BLUE
        hrun.font.name = "Calibri"

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        for cell, label, align in (
            (hdr[0], "Rule", WD_ALIGN_PARAGRAPH.LEFT),
            (hdr[1], "Source", WD_ALIGN_PARAGRAPH.LEFT),
            (hdr[2], "Severity", WD_ALIGN_PARAGRAPH.CENTER),
        ):
            shade_cell(cell, "003391")
            set_cell_text(cell, label, bold=True, color=WHITE, size=10,
                          align=align)

        for r in rules:
            cells = table.add_row().cells
            set_cell_text(cells[0], r["rule_text"].strip(), color=DARK, size=9)
            set_cell_text(cells[1], r["source"].strip(), color=GREY, size=8)
            sev = r["severity"].strip()
            shade_cell(cells[2], SEV_COLOR.get(sev, "7F8C8D"))
            set_cell_text(cells[2], sev.upper(), bold=True, color=WHITE,
                          size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

        set_col_widths(table, widths)

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH} ({total} rules across {len(ordered)} categories)")


if __name__ == "__main__":
    build()
