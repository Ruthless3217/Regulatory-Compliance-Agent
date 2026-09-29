"""Compare export artifacts — deterministic, immutable, and audit-ready.

Produces downloadable outputs:
* ``audit-snapshot.json`` — immutable, machine-readable JSON snapshot.
* ``audit-report.pdf`` — comprehensive human-readable comparison report with executive summary, metadata, SHA-256 hashes, and structured changes ledger.
* ``changes-report.docx`` — Word table report with semantic classifications.
* ``old-highlighted.pdf`` / ``new-highlighted.pdf`` — rendered pages with overlay boxes burned in.
* ``side-by-side.pdf`` — old and new pages composited per sheet.
* ``bundle.zip`` — all reports, JSON snapshots, rendered PDFs, and original inputs.
"""
import io
import os
import zipfile
import logging
from datetime import datetime, timezone
from typing import List, Optional, Dict, Any

from PIL import Image, ImageDraw
import img2pdf
from docx import Document

from app.services.render_orchestrator import renders_dir
from app.services.comparison_snapshot_service import (
    build_comparison_snapshot,
    export_snapshot_json,
    SEMANTIC_LABEL_MAP,
)

logger = logging.getLogger(__name__)

_KIND_LABEL = {
    "removed": "Removed",
    "added": "Added",
    "modified": "Modified",
    "moved": "Moved",
}

# RGBA fills / outlines for page-overlay burn-in (match the viewer palette).
_FILL = {
    "removed": (220, 38, 38, 90),   # sev-critical red
    "added": (22, 163, 74, 90),     # success green
    "moved": (124, 58, 237, 90),    # violet-600
    "modified": (14, 165, 233, 90), # sky blue
}
_OUTLINE = {
    "removed": (220, 38, 38, 255),
    "added": (22, 163, 74, 255),
    "moved": (124, 58, 237, 255),
    "modified": (14, 165, 233, 255),
}


def _page_image_path(comparison_id: str, side: str, n: int) -> str:
    return os.path.join(renders_dir(comparison_id), side, f"page-{n:04d}.png")


def _open_page(comparison_id: str, side: str, n: int) -> Optional[Image.Image]:
    path = _page_image_path(comparison_id, side, n)
    return Image.open(path) if os.path.exists(path) else None


def derive_changes(comparison) -> List[dict]:
    """Normalize either render_result.changes or diff blocks into report rows.

    Returns ``[{change_id, kind, change_type, semantic_label, structure, old_text, new_text}]``
    in canonical document order.
    """
    rr = comparison.render_result
    if comparison.render_status == "completed" and rr and rr.get("changes"):
        out: List[dict] = []
        for c in rr["changes"]:
            meta = c.get("metadata") or {}
            ctype = c.get("change_type") or c.get("semantic_type") or meta.get("category")
            if not ctype:
                ctype = {
                    "removed": "deletion",
                    "added": "insertion",
                    "moved": "reordered",
                }.get(c.get("kind"), "replacement")

            struct = dict(c.get("structure") or {})
            if c.get("section_anchor") and not struct.get("title"):
                struct["title"] = c["section_anchor"]

            out.append({
                "change_id": c["id"],
                "kind": c["kind"],
                "change_type": ctype,
                "semantic_label": SEMANTIC_LABEL_MAP.get(ctype, ctype.title()),
                "structure": struct,
                "old_text": (c.get("old") or {}).get("text", ""),
                "new_text": (c.get("new") or {}).get("text", ""),
                "old_page": (c.get("old") or {}).get("page"),
                "new_page": (c.get("new") or {}).get("page"),
                "locations": (c.get("new") or {}).get("locations") or (c.get("old") or {}).get("locations") or [],
            })
        return out

    out = []
    for i, b in enumerate(comparison.diff_result or []):
        t = b.get("type")
        if t == "equal":
            continue
        is_moved = bool(b.get("moved"))
        kind = "moved" if is_moved else {"delete": "removed", "insert": "added"}.get(t, "modified")
        ctype = "reordered" if is_moved else {"delete": "deletion", "insert": "insertion"}.get(t, "replacement")

        if t == "delete":
            old_text, new_text = b.get("old_text", ""), ""
        elif t == "insert":
            old_text, new_text = "", b.get("new_text", "")
        elif t == "replace":
            old_text = " ".join(w["text"] for w in b.get("old_words", []) if w.get("changed"))
            new_text = " ".join(w["text"] for w in b.get("new_words", []) if w.get("changed"))
        else:
            old_text = new_text = ""
        out.append({
            "change_id": f"b{i}",
            "kind": kind,
            "change_type": ctype,
            "semantic_label": SEMANTIC_LABEL_MAP.get(ctype, ctype.title()),
            "structure": {},
            "old_text": old_text,
            "new_text": new_text,
            "locations": [],
        })
    return out


def _load_font(size: int):
    """A TrueType font at ``size`` if one is on the box, else PIL's bitmap default."""
    from PIL import ImageFont
    for path in (
        "DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "arial.ttf",
    ):
        try:
            return ImageFont.truetype(path, size)
        except Exception:  # noqa: BLE001
            continue
    return ImageFont.load_default()


def _numbered_annotations(comparison, annotations):
    """Number the changes that carry a reviewer note/tags, in report order."""
    ann_by = {a.change_id: a for a in (annotations or [])}
    number_by: dict = {}
    rows: List[dict] = []
    n = 0
    for c in derive_changes(comparison):
        a = ann_by.get(c["change_id"])
        note = (getattr(a, "note", "") or "") if a else ""
        tags = list(getattr(a, "tags", []) or []) if a else []
        if note or tags:
            n += 1
            number_by[c["change_id"]] = n
            rows.append({
                "n": n,
                "kind": c["kind"],
                "change_type": c.get("change_type"),
                "semantic_label": c.get("semantic_label"),
                "structure": c.get("structure") or {},
                "note": note,
                "tags": tags,
                "old_text": c.get("old_text", ""),
                "new_text": c.get("new_text", ""),
            })
    return number_by, rows


def _draw_badge(draw, cx: float, cy: float, num: int, kind: str) -> None:
    """A small filled number badge, centred at (cx, cy), coloured by change kind."""
    r = 11
    color = _OUTLINE.get(kind, _OUTLINE["added"])
    draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=color, outline=(255, 255, 255, 255))
    font = _load_font(15)
    label = str(num)
    tw = draw.textlength(label, font=font)
    draw.text((cx - tw / 2, cy - 9), label, fill=(255, 255, 255, 255), font=font)


def audit_snapshot_json(comparison, annotations=None, filter_type: Optional[str] = None) -> bytes:
    """Generate deterministic, schema-versioned JSON audit snapshot."""
    return export_snapshot_json(comparison, annotations, filter_type).encode("utf-8")


def audit_report_pdf(comparison, annotations=None, filter_type: Optional[str] = None) -> bytes:
    """Generate comprehensive human-readable comparison report PDF."""
    snapshot = build_comparison_snapshot(comparison, annotations, filter_type)
    W, H = (1240, 1754)  # A4 at 150 DPI
    margin = 60

    title_font = _load_font(28)
    section_font = _load_font(20)
    bold_font = _load_font(16)
    regular_font = _load_font(15)
    mono_font = _load_font(13)
    small_font = _load_font(12)

    pages: List[Image.Image] = []

    def new_canvas():
        img = Image.new("RGB", (W, H), (255, 255, 255))
        return img, ImageDraw.Draw(img)

    def wrap_text(text: str, font, max_w: int, draw) -> List[str]:
        if not text:
            return [""]
        lines: List[str] = []
        for raw_line in text.splitlines():
            words = raw_line.split()
            if not words:
                lines.append("")
                continue
            cur = ""
            for word in words:
                trial = f"{cur} {word}".strip()
                if draw.textlength(trial, font=font) <= max_w or not cur:
                    cur = trial
                else:
                    lines.append(cur)
                    cur = word
            if cur:
                lines.append(cur)
        return lines or [""]

    img, draw = new_canvas()
    y = margin

    # --- Header Banner ---
    draw.rectangle([margin, y, W - margin, y + 60], fill=(15, 23, 42))  # slate-900
    draw.text((margin + 16, y + 15), "REGULATORY COMPLIANCE AGENT — COMPARISON REPORT", fill=(255, 255, 255), font=bold_font)
    y += 75

    # Title & Metadata Table
    draw.text((margin, y), snapshot["snapshot"]["title"], fill=(15, 23, 42), font=title_font)
    y += 40

    created = snapshot["snapshot"]["created_at"][:16].replace("T", " ")
    snap_id = snapshot["snapshot"]["snapshot_id"]
    content_sha = snapshot["snapshot"].get("snapshot_content_sha256", "N/A")
    scope_info = snapshot["snapshot"].get("export_scope", {})
    scope_label = f"Filter: {scope_info.get('filter')}" if scope_info.get("type") == "filtered" else "Scope: Full"

    draw.text(
        (margin, y),
        f"Snapshot ID: {snap_id}  ·  Hash: {content_sha[:16]}...  ·  {scope_label}  ·  Generated: {created} UTC",
        fill=(100, 116, 139),
        font=small_font,
    )
    y += 30

    # Document Hashes Box
    draw.rectangle([margin, y, W - margin, y + 105], fill=(248, 250, 252), outline=(226, 232, 240))
    orig_doc = snapshot["documents"]["original"]
    rev_doc = snapshot["documents"]["revised"]

    draw.text((margin + 14, y + 12), "Document A (Original):", fill=(71, 85, 105), font=bold_font)
    draw.text((margin + 200, y + 12), f"{orig_doc['filename']} ({orig_doc['content_type'].upper()})", fill=(15, 23, 42), font=regular_font)
    draw.text((margin + 14, y + 34), f"SHA-256: {orig_doc['sha256'] or 'N/A'}", fill=(100, 116, 139), font=mono_font)

    draw.text((margin + 14, y + 58), "Document B (Revised):", fill=(71, 85, 105), font=bold_font)
    draw.text((margin + 200, y + 58), f"{rev_doc['filename']} ({rev_doc['content_type'].upper()})", fill=(15, 23, 42), font=regular_font)
    draw.text((margin + 14, y + 80), f"SHA-256: {rev_doc['sha256'] or 'N/A'}", fill=(100, 116, 139), font=mono_font)
    y += 120

    # Summary Statistics Grid
    draw.text((margin, y), "Executive Summary & Semantic Statistics", fill=(15, 23, 42), font=section_font)
    y += 30

    sum_data = snapshot["summary"]
    sem_counts = sum_data["by_semantic_type"]

    stat_cards = [
        ("Total Changes", str(sum_data["total_changes"])),
        ("Numeric", str(sem_counts.get("numeric_only", 0))),
        ("Identifier", str(sem_counts.get("identifier_only", 0))),
        ("Replacement", str(sem_counts.get("replacement", 0))),
        ("Insertion", str(sum_data["insertions"])),
        ("Deletion", str(sum_data["deletions"])),
        ("Reordered", str(sum_data["reordered"])),
        ("Formatting", str(sem_counts.get("punctuation_only", 0) + sem_counts.get("whitespace_only", 0))),
    ]

    card_w = (W - 2 * margin - 7 * 10) / 8
    for idx, (lbl, val) in enumerate(stat_cards):
        cx = margin + idx * (card_w + 10)
        draw.rectangle([cx, y, cx + card_w, y + 50], fill=(241, 245, 249), outline=(203, 213, 225))
        tw_v = draw.textlength(val, font=bold_font)
        draw.text((cx + (card_w - tw_v) / 2, y + 8), val, fill=(15, 23, 42), font=bold_font)
        tw_l = draw.textlength(lbl, font=small_font)
        draw.text((cx + (card_w - tw_l) / 2, y + 30), lbl, fill=(100, 116, 139), font=small_font)
    y += 70

    # Changes Table Section
    draw.text((margin, y), "Detailed Changes Ledger", fill=(15, 23, 42), font=section_font)
    y += 30

    # Table Header
    draw.rectangle([margin, y, W - margin, y + 26], fill=(226, 232, 240))
    draw.text((margin + 8, y + 6), "#", fill=(51, 65, 85), font=bold_font)
    draw.text((margin + 40, y + 6), "Type", fill=(51, 65, 85), font=bold_font)
    draw.text((margin + 170, y + 6), "Section / Scope", fill=(51, 65, 85), font=bold_font)
    draw.text((margin + 430, y + 6), "Original Text", fill=(51, 65, 85), font=bold_font)
    draw.text((margin + 780, y + 6), "Revised Text", fill=(51, 65, 85), font=bold_font)
    y += 28

    ann_map = {a["change_id"]: a for a in snapshot["review_audit_trail"]}

    for idx, ch in enumerate(snapshot["changes"], 1):
        c_id = ch["change_id"]
        c_type = ch.get("semantic_label") or ch.get("change_type", "Changed")
        struct_title = ch.get("structure", {}).get("title") or "—"
        old_t = ch.get("old_text") or ""
        new_t = ch.get("new_text") or ""
        ann = ann_map.get(c_id)

        # Multi-location note
        loc_count = len(ch.get("locations") or [])
        loc_str = f" ({loc_count} locs)" if loc_count > 1 else ""

        old_lines = wrap_text(old_t, regular_font, 330, draw)
        new_lines = wrap_text(new_t, regular_font, 330, draw)
        max_lines = max(len(old_lines), len(new_lines), 1)

        row_h = max_lines * 20 + (30 if ann else 14)

        if y + row_h > H - margin:
            pages.append(img)
            img, draw = new_canvas()
            y = margin
            # Draw header on next page
            draw.rectangle([margin, y, W - margin, y + 26], fill=(226, 232, 240))
            draw.text((margin + 8, y + 6), "#", fill=(51, 65, 85), font=bold_font)
            draw.text((margin + 40, y + 6), "Type", fill=(51, 65, 85), font=bold_font)
            draw.text((margin + 170, y + 6), "Section / Scope", fill=(51, 65, 85), font=bold_font)
            draw.text((margin + 430, y + 6), "Original Text", fill=(51, 65, 85), font=bold_font)
            draw.text((margin + 780, y + 6), "Revised Text", fill=(51, 65, 85), font=bold_font)
            y += 28

        # Row Background & Borders
        draw.rectangle([margin, y, W - margin, y + row_h], fill=(255, 255, 255), outline=(241, 245, 249))

        draw.text((margin + 8, y + 6), str(idx), fill=(71, 85, 105), font=regular_font)
        draw.text((margin + 40, y + 6), f"{c_type}{loc_str}", fill=(15, 23, 42), font=bold_font)

        # Section title truncated if long
        s_title_trunc = struct_title[:28] + ("…" if len(struct_title) > 28 else "")
        draw.text((margin + 170, y + 6), s_title_trunc, fill=(100, 116, 139), font=regular_font)

        # Old text lines (red tint)
        for li, ln in enumerate(old_lines):
            draw.text((margin + 430, y + 6 + li * 20), ln, fill=(185, 28, 28), font=regular_font)

        # New text lines (green tint)
        for li, ln in enumerate(new_lines):
            draw.text((margin + 780, y + 6 + li * 20), ln, fill=(21, 128, 61), font=regular_font)

        # Reviewer annotation note if present
        if ann:
            note_y = y + max_lines * 20 + 6
            note_str = f"Reviewer Note: {ann.get('note') or '—'}"
            if ann.get("tags"):
                note_str += f"  [Tags: {', '.join(ann['tags'])}]"
            draw.text((margin + 40, note_y), note_str, fill=(79, 70, 229), font=small_font)

        y += row_h + 4

    pages.append(img)

    # Convert all rendered PNG pages to single PDF
    out_pdf_bytes: List[bytes] = []
    for p in pages:
        buf = io.BytesIO()
        p.save(buf, format="PNG", dpi=(150, 150))
        out_pdf_bytes.append(buf.getvalue())

    return img2pdf.convert(out_pdf_bytes)


def changes_report_docx(comparison, annotations) -> bytes:
    """The change list as a Word document with semantic classifications."""
    ann_by_change = {a.change_id: a for a in (annotations or [])}
    changes = derive_changes(comparison)

    doc = Document()
    doc.add_heading(comparison.title or "Document comparison", level=0)

    created = comparison.created_at.strftime("%Y-%m-%d %H:%M") if comparison.created_at else "—"
    meta = doc.add_paragraph()
    meta.add_run(
        f"Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC · Created {created} · Engine v1.0"
    ).italic = True

    counts: dict = {}
    for c in changes:
        lbl = c.get("semantic_label") or _KIND_LABEL.get(c["kind"], c["kind"])
        counts[lbl] = counts.get(lbl, 0) + 1

    summary = doc.add_paragraph()
    summary.add_run("Semantic Summary: ").bold = True
    summary.add_run(
        ", ".join(f"{k}: {counts[k]}" for k in sorted(counts)) or "No changes"
    )

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["#", "Type", "Section", "Original", "Revised", "Note", "Tags"]):
        cell.text = head

    for idx, c in enumerate(changes, 1):
        a = ann_by_change.get(c["change_id"])
        cells = table.add_row().cells
        cells[0].text = str(idx)
        cells[1].text = c.get("semantic_label") or _KIND_LABEL.get(c["kind"], c["kind"])
        cells[2].text = c.get("structure", {}).get("title") or "—"
        cells[3].text = c["old_text"] or ""
        cells[4].text = c["new_text"] or ""
        cells[5].text = (a.note or "") if a else ""
        cells[6].text = ", ".join(a.tags) if a and a.tags else ""

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def highlighted_pdf(comparison, side: str, annotations=None) -> bytes:
    """One side's rendered pages with overlay boxes burned in, as a PDF."""
    rr = comparison.render_result or {}
    kind_by_change = {c["id"]: c["kind"] for c in rr.get("changes", [])}
    pages = (rr.get(side) or {}).get("pages", [])
    number_by, rows = _numbered_annotations(comparison, annotations)

    png_pages: List[bytes] = []
    for pg in pages:
        img = _open_page(str(comparison.id), side, pg["n"])
        if img is None:
            continue
        img = img.convert("RGBA")
        w_px, h_px = img.size
        sx = w_px / pg["w_pt"] if pg.get("w_pt") else 1.0
        sy = h_px / pg["h_pt"] if pg.get("h_pt") else 1.0
        overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(overlay)
        for b in pg.get("boxes", []):
            kind = "moved" if kind_by_change.get(b["change_id"]) == "moved" else b["type"]
            draw.rectangle(
                [b["x0"] * sx, b["y0"] * sy, b["x1"] * sx, b["y1"] * sy],
                fill=_FILL.get(kind, _FILL["added"]),
                outline=_OUTLINE.get(kind, _OUTLINE["added"]),
                width=1,
            )
            num = number_by.get(b["change_id"])
            if num is not None:
                _draw_badge(draw, max(11.0, b["x0"] * sx), max(11.0, b["y0"] * sy), num, kind)
        flat = Image.alpha_composite(img, overlay).convert("RGB")
        out = io.BytesIO()
        flat.save(out, format="PNG")
        png_pages.append(out.getvalue())

    if not png_pages:
        raise ValueError("No rendered pages available to highlight")
    if rows:
        png_pages.extend(_comments_pages(rows))
    return img2pdf.convert(png_pages)


def side_by_side_pdf(comparison, annotations=None) -> bytes:
    """Old and new pages composited left/right per sheet, as a PDF."""
    rr = comparison.render_result or {}
    old_pages = (rr.get("old") or {}).get("pages", [])
    new_pages = (rr.get("new") or {}).get("pages", [])
    gap = 24
    sheets: List[bytes] = []

    for k in range(max(len(old_pages), len(new_pages))):
        left = _open_page(str(comparison.id), "old", old_pages[k]["n"]) if k < len(old_pages) else None
        right = _open_page(str(comparison.id), "new", new_pages[k]["n"]) if k < len(new_pages) else None
        if left is None and right is None:
            continue
        lw, lh = (left.size if left else (0, 0))
        rw, rh = (right.size if right else (0, 0))
        if left is None:
            lw = rw
        if right is None:
            rw = lw
        canvas = Image.new("RGB", (lw + gap + rw, max(lh, rh) or 1), (255, 255, 255))
        if left is not None:
            canvas.paste(left.convert("RGB"), (0, 0))
        if right is not None:
            canvas.paste(right.convert("RGB"), (lw + gap, 0))
        out = io.BytesIO()
        canvas.save(out, format="PNG")
        sheets.append(out.getvalue())

    if not sheets:
        raise ValueError("No rendered pages available")
    _, rows = _numbered_annotations(comparison, annotations)
    if rows:
        sheets.extend(_comments_pages(rows))
    return img2pdf.convert(sheets)


def bundle_zip(comparison, annotations) -> bytes:
    """Report + JSON Snapshot + PDFs + original inputs, zipped."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("audit-snapshot.json", audit_snapshot_json(comparison, annotations))
        z.writestr("audit-report.pdf", audit_report_pdf(comparison, annotations))
        z.writestr("changes-report.docx", changes_report_docx(comparison, annotations))
        if comparison.render_status == "completed" and comparison.render_result:
            try:
                z.writestr("old-highlighted.pdf", highlighted_pdf(comparison, "old", annotations))
                z.writestr("new-highlighted.pdf", highlighted_pdf(comparison, "new", annotations))
                z.writestr("side-by-side.pdf", side_by_side_pdf(comparison, annotations))
            except Exception as e:
                logger.warning("bundle_zip: skipping visual PDFs for %s: %s", comparison.id, e)
        for label, path in (("original", comparison.old_file_path), ("revised", comparison.new_file_path)):
            if path and os.path.exists(path):
                z.write(path, arcname=f"inputs/{label}-{os.path.basename(path)}")
    return buf.getvalue()
