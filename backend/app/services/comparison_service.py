"""
Document comparison service — paragraph extraction and word-level diffing.

No LLM calls; pure text processing (python-docx / pdfplumber for extraction,
stdlib difflib for diffing), so comparisons run synchronously in the API layer.

Cross-format normalization
--------------------------
Two versions of a document are frequently *different file types* — e.g. a DOCX
first draft vs. a finalized PDF. Raw paragraph strings from the two extractors
never match byte-for-byte (PDF collapses paragraphs into page-sized blobs, adds
running headers/footers, wraps lines, and the DOCX is a template full of
``<placeholder>`` fields). Diffing those raw chunks marks the whole document as
removed/added.

To fix this we compare *sentence-level segments* aligned on a *normalized key*
(whitespace/case/heading-marker-insensitive), and treat a segment whose only
difference is a filled-in template placeholder as unchanged.
"""
import logging
import re
import string
from collections import Counter
from difflib import SequenceMatcher
from typing import List, Optional, Set

logger = logging.getLogger(__name__)


def split_text_paragraphs(text: str) -> List[str]:
    """Split raw text into paragraphs on one-or-more blank lines."""
    if not text or not text.strip():
        return []
    blocks = re.split(r"\n\s*\n", text.strip())
    return [b.strip() for b in blocks if b.strip()]


def extract_docx_paragraphs(file_path: str) -> List[str]:
    """Extract paragraph text from a DOCX file, one entry per Word paragraph."""
    from docx import Document
    doc = Document(file_path)
    paragraphs: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = getattr(para.style, "name", "") or ""
        if style.startswith("Heading") or style == "Title":
            paragraphs.append(f"## {text}")
        else:
            paragraphs.append(text)
    return paragraphs


def extract_pdf_paragraphs(file_path: str) -> List[str]:
    """Extract paragraph text from a PDF: page text joined, then split on blank lines."""
    import pdfplumber
    pages: List[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                pages.append(page_text)
    return split_text_paragraphs("\n\n".join(pages))


def extract_paragraphs(
    file_path: Optional[str], content_type: str, pasted_text: Optional[str] = None
) -> List[str]:
    """Dispatch extraction by content_type: docx/pdf read from file_path, else split pasted_text."""
    if content_type == "docx":
        if not file_path:
            raise ValueError("docx content_type requires file_path")
        return extract_docx_paragraphs(file_path)
    if content_type == "pdf":
        if not file_path:
            raise ValueError("pdf content_type requires file_path")
        return extract_pdf_paragraphs(file_path)
    if file_path:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return split_text_paragraphs(f.read())
    return split_text_paragraphs(pasted_text or "")


# ---------------------------------------------------------------------------
# Cross-format normalization: sentence segmentation, header/footer stripping,
# placeholder-aware matching keys.
# ---------------------------------------------------------------------------

# Split on whitespace that follows a sentence terminator. Sentences survive
# reflowing between formats far better than paragraphs or physical lines.
_SENTENCE_BOUNDARY = re.compile(r"(?<=[.:;!?])\s+")

# Heading markers we add to DOCX headings (## ), so both sides key the same.
_HEADING_MARKER = re.compile(r"^#{1,6}\s*")

# A line that is just a page number ("3", "Page 3", "3 of 21", "3/21").
_PAGE_NUMBER = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.IGNORECASE)

# Hyphen introduced by a physical line-wrap: "insur-\nance" -> "insurance".
_HYPHEN_WRAP = re.compile(r"(\w)-\n(\w)")

# Word tokenizer: keep an angle-bracket template field (which may contain spaces,
# e.g. "<Name of the Policyholder>") as ONE token, absorbing any punctuation that
# trails it (a terminating "." etc.); otherwise split on whitespace.
_TOKEN_RE = re.compile(r"<[^>\n]*>[^\s<]*|\S+")

# A whole token that is a template placeholder standing in for a value: an
# angle-bracket field, a run of underscores, or a run of X's. Filling it in is an
# expected edit, not a content change.
_PLACEHOLDER_TOKEN = re.compile(r"<[^>\n]*>|_{2,}|[Xx]{2,}")


def split_sentences(text: str) -> List[str]:
    """Split a block of text into trimmed, non-empty sentence-level segments."""
    text = (text or "").strip()
    if not text:
        return []
    return [s.strip() for s in _SENTENCE_BOUNDARY.split(text) if s.strip()]


def normalize_for_match(text: str) -> str:
    """Normalized alignment key: drop heading markers, collapse whitespace, lowercase.

    Only used to decide whether two segments are 'the same'; the original text is
    always what gets displayed.
    """
    t = _HEADING_MARKER.sub("", (text or "").strip())
    t = re.sub(r"\s+", " ", t)
    return t.strip().lower()


def _detect_running_lines(page_lines: List[List[str]]) -> Set[str]:
    """Find lines that repeat in the top/bottom band of most pages (headers/footers).

    Needs at least 3 pages to distinguish a running header from ordinary content.
    """
    if len(page_lines) < 3:
        return set()
    band = Counter()
    for lines in page_lines:
        for ln in lines[:3] + lines[-3:]:
            band[ln] += 1
    threshold = max(3, len(page_lines) // 2)
    return {ln for ln, count in band.items() if count >= threshold}


def _pdf_text_layer_lines(file_path: str) -> List[List[str]]:
    """Per-page non-empty text lines from a PDF's selectable text layer."""
    import pdfplumber
    page_lines: List[List[str]] = []
    with pdfplumber.open(file_path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            # A single malformed page (bad font map, broken content stream) must
            # not fail the whole comparison — skip it and keep going.
            try:
                txt = page.extract_text() or ""
            except Exception as e:  # noqa: BLE001
                logger.warning("PDF page %s text extraction failed, skipping: %s", page_no, e)
                txt = ""
            page_lines.append([ln.strip() for ln in txt.split("\n") if ln.strip()])
    return page_lines


def _pdf_ocr_lines(file_path: str) -> List[List[str]]:
    """OCR a scanned/image PDF into per-page text lines.

    Rasterizes each page with pypdfium2 (already a dependency for the pixel view)
    and runs Tesseract via pytesseract. Returns ``[]`` — never raises — if OCR is
    disabled or the tooling isn't installed, so the caller can fall back to the
    clear "run OCR" message. Capped at ``settings.compare_ocr_page_cap`` pages.
    """
    from app.config import settings
    if not settings.compare_ocr_enabled:
        return []
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except Exception as e:  # noqa: BLE001 — OCR tooling not present in this image
        logger.warning("OCR fallback unavailable (%s); scanned PDF will not be read", e)
        return []

    page_lines: List[List[str]] = []
    try:
        pdf = pdfium.PdfDocument(file_path)
    except Exception as e:  # noqa: BLE001
        logger.warning("OCR: could not open PDF %s: %s", file_path, e)
        return []
    try:
        cap = min(len(pdf), max(1, settings.compare_ocr_page_cap))
        for i in range(cap):
            try:
                # scale 3.0 ≈ 216 DPI — enough for Tesseract without huge bitmaps.
                # draw_annots=False: reviewer annotations (FreeText notes, stamps,
                # popups) must not be rasterized into text that downstream
                # consumers treat as page content — pypdfium2 draws them by default.
                pil = pdf[i].render(scale=3.0, draw_annots=False).to_pil()
                txt = pytesseract.image_to_string(pil) or ""
            except Exception as e:  # noqa: BLE001 — skip a page OCR can't handle
                logger.warning("OCR failed on page %s: %s", i + 1, e)
                txt = ""
            page_lines.append([ln.strip() for ln in txt.split("\n") if ln.strip()])
    finally:
        pdf.close()
    if any(page_lines):
        logger.info("OCR fallback extracted text from %s (%d pages)", file_path, len(page_lines))
    return page_lines


def extract_pdf_segments(file_path: str) -> List[str]:
    """Extract sentence-level segments from a PDF, normalized for cross-format diff.

    Strips repeated running headers/footers and page numbers, de-hyphenates
    line-wraps, unwraps physical lines back into flowing text, then segments into
    sentences. Falls back to OCR for a scanned/image PDF (no text layer); raises
    ValueError only if neither the text layer nor OCR yields any text.
    """
    page_lines = _pdf_text_layer_lines(file_path)

    if not any(page_lines):  # no selectable text — try OCR before giving up
        page_lines = _pdf_ocr_lines(file_path)

    if not any(page_lines):
        raise ValueError(
            "No extractable text found in the PDF — it looks like a scanned image "
            "and OCR could not read it. Provide a text-based PDF or run OCR first."
        )

    running = _detect_running_lines(page_lines)
    cleaned: List[str] = []
    for lines in page_lines:
        body = [ln for ln in lines if ln not in running and not _PAGE_NUMBER.match(ln)]
        text = "\n".join(body)
        text = _HYPHEN_WRAP.sub(r"\1\2", text)  # join hyphenated line-wraps
        text = text.replace("\n", " ")          # unwrap remaining physical lines
        cleaned.append(text)
    return split_sentences(" ".join(cleaned))


def _docx_paragraph_text(p_el) -> str:
    """Full visible text of a ``<w:p>`` element.

    python-docx's ``Paragraph.text`` only concatenates runs that are *direct*
    ``<w:r>`` children of the paragraph, so it silently drops any run wrapped in
    another element — most importantly tracked-change insertions (``<w:ins>``),
    moved-in runs (``<w:moveTo>``) and hyperlinks. In a reviewed document those
    can be a large share of the body, which is exactly the content that then goes
    "missing" from a comparison. We instead walk every text node under the
    paragraph in document order.

    Excluded on purpose: tracked *deletions* (which use ``<w:delText>``, not
    ``<w:t>``, so they never appear here), the old location of a moved run
    (``<w:moveFrom>`` — its twin ``<w:moveTo>`` carries the surviving copy), and
    any text-box content (``<w:txbxContent>``) anchored inside the paragraph.
    """
    from docx.oxml.ns import qn

    t_tag, tab_tag = qn("w:t"), qn("w:tab")
    br_tag, cr_tag = qn("w:br"), qn("w:cr")
    move_from_tag, txbx_tag = qn("w:moveFrom"), qn("w:txbxContent")

    parts: List[str] = []
    for node in p_el.iter(t_tag, tab_tag, br_tag, cr_tag):
        anc = node.getparent()
        skip = False
        while anc is not None and anc is not p_el:
            if anc.tag == move_from_tag or anc.tag == txbx_tag:
                skip = True
                break
            anc = anc.getparent()
        if skip:
            continue
        if node.tag == t_tag:
            parts.append(node.text or "")
        elif node.tag == tab_tag:
            parts.append("\t")
        else:  # w:br / w:cr
            parts.append("\n")
    return "".join(parts)


def _p_is_heading(p_el, doc) -> bool:
    """True if a ``<w:p>`` element carries a Heading/Title paragraph style."""
    from docx.text.paragraph import Paragraph
    try:
        name = getattr(Paragraph(p_el, doc).style, "name", "") or ""
    except Exception:
        name = ""
    return name.startswith("Heading") or name == "Title"


def _iter_docx_block_texts(doc) -> List[str]:
    """Every block of body text in document order — nothing structurally dropped.

    A single recursive walk over the XML tree captures text wherever Word can put
    it, because ``doc.paragraphs`` / ``doc.tables`` only see the top layer:

    * **Table cells** — walked by raw ``<w:tc>`` iteration rather than
      python-docx's ``row.cells``. That both avoids python-docx's grid expansion
      (which duplicated vertically-merged cells) and sidesteps the ``id(cell._tc)``
      de-dup trap (ephemeral lxml proxies whose ``id()`` CPython recycles, so
      distinct cells collided and were skipped, dropping real content). Each
      ``<w:tc>`` is a distinct element, so merges need no de-duping: a
      vertically-merged continuation cell is its own (empty) ``<w:tc>``.
    * **Content controls** (``<w:sdt>``) — descend into ``<w:sdtContent>`` at
      block level and inside table rows; inline controls are covered by
      :func:`_docx_paragraph_text`.
    * **Text boxes** (``<w:txbxContent>``) — emitted as their own blocks, since
      their text is anchored inside a run's drawing, invisible to a paragraph walk.

    Tracked-change and hyperlink handling lives in :func:`_docx_paragraph_text`.
    """
    from docx.oxml.ns import qn

    P, TBL, SDT, SDTC = qn("w:p"), qn("w:tbl"), qn("w:sdt"), qn("w:sdtContent")
    TR, TC, TXBX = qn("w:tr"), qn("w:tc"), qn("w:txbxContent")

    out: List[str] = []

    def emit_p(p_el) -> None:
        text = _docx_paragraph_text(p_el).strip()
        if text:
            out.append(f"## {text}" if _p_is_heading(p_el, doc) else text)
        # Text boxes anchored in this paragraph become their own blocks, in place.
        for txbx in p_el.iter(TXBX):
            walk(txbx)

    def walk_row(tr) -> None:
        for cellish in tr.iterchildren():
            if cellish.tag == TC:
                walk(cellish)
            elif cellish.tag == SDT:  # a content-control-wrapped cell
                content = cellish.find(SDTC)
                if content is not None:
                    for tc in content.iterchildren(TC):
                        walk(tc)

    def walk_table(tbl_el) -> None:
        for child in tbl_el.iterchildren():
            if child.tag == TR:
                walk_row(child)
            elif child.tag == SDT:  # content-control-wrapped row(s)
                content = child.find(SDTC)
                if content is not None:
                    for tr in content.iterchildren(TR):
                        walk_row(tr)

    def walk(container) -> None:
        for child in container.iterchildren():
            tag = child.tag
            if tag == P:
                emit_p(child)
            elif tag == TBL:
                walk_table(child)
            elif tag == SDT:
                content = child.find(SDTC)
                if content is not None:
                    walk(content)

    walk(doc.element.body)
    return out


def extract_docx_segments(file_path: str) -> List[str]:
    """Extract sentence-level segments from a DOCX (headings kept whole, tables included)."""
    from docx import Document
    doc = Document(file_path)
    segments: List[str] = []
    for block_text in _iter_docx_block_texts(doc):
        if block_text.startswith("#"):
            segments.append(block_text)
        else:
            segments.extend(split_sentences(block_text))
    return segments


def extract_segments(
    file_path: Optional[str], content_type: str, pasted_text: Optional[str] = None
) -> List[str]:
    """Dispatch sentence-level extraction by content_type (docx/pdf/text)."""
    if content_type == "docx":
        if not file_path:
            raise ValueError("docx content_type requires file_path")
        return extract_docx_segments(file_path)
    if content_type == "pdf":
        if not file_path:
            raise ValueError("pdf content_type requires file_path")
        return extract_pdf_segments(file_path)
    if file_path:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            raw = f.read()
    else:
        raw = pasted_text or ""
    segments: List[str] = []
    for para in split_text_paragraphs(raw):
        segments.extend(split_sentences(para))
    return segments


def _friendly_extract_message(content_type: str, exc: Exception) -> str:
    """Turn a raw extraction exception into a message a reviewer can act on."""
    name = type(exc).__name__
    raw = str(exc).strip()
    if content_type == "docx":
        if name in ("PackageNotFoundError", "BadZipFile") or "not a Word file" in raw or "Package not found" in raw:
            return (
                "this doesn't open as a .docx. If it's an older .doc file, open it in "
                "Word and use Save As → Word Document (.docx), then upload again."
            )
        return f"the Word file could not be read ({raw or name})."
    if content_type == "pdf":
        low = raw.lower()
        if "password" in low or "encrypt" in low:
            return "the PDF is password-protected — remove the protection and re-upload."
        return f"the PDF could not be parsed ({raw or name})."
    return raw or name


def extract_segments_labeled(
    file_path: Optional[str], content_type: str, pasted_text: Optional[str], side_label: str
) -> List[str]:
    """extract_segments with a clear, side-labelled error on failure.

    ValueErrors raised deliberately downstream (e.g. the scanned-PDF hint) are
    already user-friendly and pass through unchanged; anything else is wrapped so
    the reviewer sees which side failed and why, instead of a raw library trace.
    """
    try:
        return extract_segments(file_path, content_type, pasted_text)
    except ValueError:
        raise
    except Exception as e:  # noqa: BLE001 — translate to a reviewer-facing message
        logger.error("Extraction failed for %s side (%s): %s", side_label, content_type, e, exc_info=True)
        raise ValueError(f"Couldn't read the {side_label} document: {_friendly_extract_message(content_type, e)}")


# ---------------------------------------------------------------------------
# Word-token alignment
# ---------------------------------------------------------------------------
# Cross-format comparison aligns at the *word-token* level over a flat stream
# rather than sentence-by-sentence, so identical prose stays aligned even when
# the two extractors chunk it differently — the failure mode that otherwise marks
# a whole document as removed + re-added. The token opcodes are then re-grouped
# into sentence-sized rows for a readable two-column display.

_PUNCT = string.punctuation
_SENTENCE_TERMINATORS = ".:;!?"
_CLOSERS = "\"')]}"
# Punctuation that may bracket a placeholder without being part of it, e.g. the
# terminator in "<dd/mm/yyyy>." or "______." — stripped before detection.
_PH_TRIM = "\"'([{)]}.,;:!?"


def _is_placeholder(tok: str) -> bool:
    """True if the token (bare of any surrounding punctuation) is a template
    placeholder: an angle-bracket field, an underscore blank, or a run of X's."""
    return bool(_PLACEHOLDER_TOKEN.fullmatch(tok.strip(_PH_TRIM)))


def _norm_token(tok: str) -> str:
    """Matching key for one token: template placeholders collapse to a shared
    wildcard; everything else is lowercased with surrounding punctuation dropped."""
    if _is_placeholder(tok):
        return "\x00ph"
    stripped = tok.strip(_PUNCT).lower()
    return stripped or tok.lower()


def _tokenize(segments: List[str]) -> List[tuple]:
    """Flatten segments into a flat list of (display_text, match_key) word tokens.

    Heading markers ('## ') are stripped so a heading keys like its own prose.
    """
    tokens: List[tuple] = []
    for seg in segments:
        seg = _HEADING_MARKER.sub("", seg)
        for m in _TOKEN_RE.finditer(seg):
            t = m.group(0)
            tokens.append((t, _norm_token(t)))
    return tokens


def _ends_sentence(text: str) -> bool:
    core = text.rstrip(_CLOSERS)
    return bool(core) and core[-1] in _SENTENCE_TERMINATORS


def _join(words: List[dict]) -> str:
    return " ".join(w["text"] for w in words)


def _split_words_into_sentences(words: List[dict]) -> List[List[dict]]:
    """Break a run of word dicts into sentence-sized chunks for readable rows."""
    out: List[List[dict]] = []
    cur: List[dict] = []
    for w in words:
        cur.append(w)
        if _ends_sentence(w["text"]):
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def _suppress_placeholder_fills(old_words: List[dict], new_words: List[dict]) -> None:
    """Clear the `changed` flag on spans whose only difference is a filled-in
    template placeholder (e.g. '<XX>' -> '10'): an expected fill, not a real edit.

    Mutates the word dicts in place.
    """
    o_norm = [_norm_token(w["text"]) for w in old_words]
    n_norm = [_norm_token(w["text"]) for w in new_words]
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, o_norm, n_norm, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        o_span, n_span = old_words[i1:i2], new_words[j1:j2]
        old_is_ph = bool(o_span) and all(_is_placeholder(w["text"]) for w in o_span)
        new_is_ph = bool(n_span) and all(_is_placeholder(w["text"]) for w in n_span)
        if old_is_ph or new_is_ph:
            for w in o_span + n_span:
                w["changed"] = False


def _emit_blocks(old_words: List[dict], new_words: List[dict]) -> List[dict]:
    """Turn one accumulated row buffer into display block(s)."""
    if old_words and not new_words:
        return [{"type": "delete", "old_text": _join(old_words)}]
    if new_words and not old_words:
        return [{"type": "insert", "new_text": _join(new_words)}]

    _suppress_placeholder_fills(old_words, new_words)

    if not any(w["changed"] for w in old_words) and not any(w["changed"] for w in new_words):
        return [{"type": "equal", "old_text": _join(old_words), "new_text": _join(new_words)}]

    # No unchanged tokens on either side => the two runs are unrelated content,
    # not a modification of one another; show them as separate removals and
    # additions rather than forcing them to face each other in one row.
    if all(w["changed"] for w in old_words) and all(w["changed"] for w in new_words):
        return (
            [{"type": "delete", "old_text": _join(s)} for s in _split_words_into_sentences(old_words)]
            + [{"type": "insert", "new_text": _join(s)} for s in _split_words_into_sentences(new_words)]
        )

    return [{"type": "replace", "old_words": old_words, "new_words": new_words}]


# ---------------------------------------------------------------------------
# Move detection (post-pass)
# ---------------------------------------------------------------------------
# A run deleted from one place and inserted verbatim elsewhere is a *move*, not
# an unrelated delete + add. We tag both ends so the UI can link/dim them instead
# of shouting a removal and an addition. Matching is on the normalized token
# stream (the same key used for diffing) with a minimum run length so trivial
# repeated phrases don't pair up.

_MOVE_MIN_TOKENS = 5


def _move_key(text: str) -> tuple:
    """Normalized token tuple for move matching (shares the diff normalizer)."""
    stripped = _HEADING_MARKER.sub("", text or "")
    return tuple(_norm_token(m.group(0)) for m in _TOKEN_RE.finditer(stripped))


def detect_moves_in_blocks(blocks: List[dict]) -> List[dict]:
    """Tag delete/insert block pairs with identical normalized text as moves.

    Mutates and returns ``blocks``: each paired delete and insert gains
    ``moved=True`` and a shared ``move_id`` ("m0", "m1", ...). Only ``delete`` and
    ``insert`` blocks are candidates (a ``replace`` is an in-place edit, not a
    move); pairing is 1:1. No pair found => blocks are returned unchanged.
    """
    deletes: List[tuple] = []   # (key, block_index)
    inserts: List[tuple] = []
    for i, b in enumerate(blocks):
        if b["type"] == "delete":
            key = _move_key(b["old_text"])
            if len(key) >= _MOVE_MIN_TOKENS:
                deletes.append((key, i))
        elif b["type"] == "insert":
            key = _move_key(b["new_text"])
            if len(key) >= _MOVE_MIN_TOKENS:
                inserts.append((key, i))

    used_ins: set = set()
    move_no = 0
    for dkey, di in deletes:
        for pos, (ikey, ii) in enumerate(inserts):
            if pos in used_ins or ikey != dkey:
                continue
            used_ins.add(pos)
            mid = f"m{move_no}"
            move_no += 1
            blocks[di]["moved"] = True
            blocks[di]["move_id"] = mid
            blocks[ii]["moved"] = True
            blocks[ii]["move_id"] = mid
            break
    return blocks


def _detect_moves_in_changes(new_marks: List[dict], changes: List[dict]) -> None:
    """Merge removed/added change pairs with identical normalized text into one
    ``moved`` change (kind carries both old_text and new_text).

    Mutates ``changes`` in place and relabels the added side's ``new_marks`` onto
    the surviving (removed) change id so a downstream renderer can attach BOTH an
    old-side and new-side box to the single moved change.
    """
    removed = [(c, _move_key(c["old_text"])) for c in changes if c["kind"] == "removed"]
    added = [(c, _move_key(c["new_text"])) for c in changes if c["kind"] == "added"]
    used_added: set = set()
    drop_ids: set = set()
    for rc, rkey in removed:
        if len(rkey) < _MOVE_MIN_TOKENS:
            continue
        for ac, akey in added:
            if ac["id"] in used_added or akey != rkey:
                continue
            used_added.add(ac["id"])
            for m in new_marks:
                if m["change_id"] == ac["id"]:
                    m["change_id"] = rc["id"]
            rc["kind"] = "moved"
            rc["new_text"] = ac["new_text"]
            drop_ids.add(ac["id"])
            break
    if drop_ids:
        changes[:] = [c for c in changes if c["id"] not in drop_ids]


def build_diff(old_paragraphs: List[str], new_paragraphs: List[str]) -> List[dict]:
    """Align two documents at the word-token level and return ordered diff blocks.

    Both sides are flattened to a normalized token stream (placeholder-aware,
    heading-marker/whitespace/case-insensitive) and aligned with difflib. Because
    alignment ignores how each extractor chunked the text, identical content stays
    matched across formats even when segmentation drifts. The resulting token
    opcodes are re-grouped into sentence-sized rows carrying the existing
    equal/delete/insert/replace block schema the frontend already renders.
    """
    old_tokens = _tokenize(old_paragraphs)
    new_tokens = _tokenize(new_paragraphs)
    matcher = SequenceMatcher(
        None, [t[1] for t in old_tokens], [t[1] for t in new_tokens], autojunk=False
    )

    blocks: List[dict] = []
    cur_old: List[dict] = []
    cur_new: List[dict] = []

    def flush() -> None:
        nonlocal cur_old, cur_new
        if cur_old or cur_new:
            blocks.extend(_emit_blocks(cur_old, cur_new))
            cur_old, cur_new = [], []

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for k in range(i2 - i1):
                ot, nt = old_tokens[i1 + k], new_tokens[j1 + k]
                cur_old.append({"text": ot[0], "changed": False})
                cur_new.append({"text": nt[0], "changed": False})
                if _ends_sentence(ot[0]) or _ends_sentence(nt[0]):
                    flush()
        elif tag == "replace":
            for ot in old_tokens[i1:i2]:
                cur_old.append({"text": ot[0], "changed": True})
            for nt in new_tokens[j1:j2]:
                cur_new.append({"text": nt[0], "changed": True})
        elif tag == "delete":
            for ot in old_tokens[i1:i2]:
                cur_old.append({"text": ot[0], "changed": True})
                if not cur_new and _ends_sentence(ot[0]):
                    flush()
        elif tag == "insert":
            for nt in new_tokens[j1:j2]:
                cur_new.append({"text": nt[0], "changed": True})
                if not cur_old and _ends_sentence(nt[0]):
                    flush()
    flush()
    return detect_moves_in_blocks(blocks)


def match_query_in_words(words: List[dict], query: str, page: int, cap: int = 200) -> List[dict]:
    """Case-insensitive substring search of ``query`` across a page's word stream.

    ``words`` are pdfplumber-style dicts (``text``/``x0``/``top``/``x1``/``bottom``);
    they are joined with single spaces and every occurrence of ``query`` maps back
    to the union bbox of the words it covers. Returns ``[{page, bbox:[x0,y0,x1,y1]}]``
    (top-left origin, PDF points), at most ``cap`` hits. Pure — no I/O — so it is
    unit-testable without a real PDF.
    """
    ql = (query or "").lower()
    if not ql:
        return []
    text = ""
    spans: List[tuple] = []  # (start_char, end_char, word)
    for w in words:
        start = len(text)
        text += w["text"]
        spans.append((start, len(text), w))
        text += " "
    tl = text.lower()

    hits: List[dict] = []
    pos = tl.find(ql)
    while pos != -1 and len(hits) < cap:
        end = pos + len(ql)
        covered = [w for (s, e, w) in spans if s < end and e > pos]
        if covered:
            hits.append({
                "page": page,
                "bbox": [
                    min(float(w["x0"]) for w in covered),
                    min(float(w["top"]) for w in covered),
                    max(float(w["x1"]) for w in covered),
                    max(float(w["bottom"]) for w in covered),
                ],
            })
        pos = tl.find(ql, pos + 1)
    return hits


def word_level_ops(old_texts: List[str], new_texts: List[str]):
    """Align two positioned-word streams by normalized text and tag each word.

    Returns (old_marks, new_marks, changes). Placeholder fills (e.g. '<XX>' filled
    with a value) are treated as equal. Used by the pixel-faithful overlay; shares
    the token normalizer with `build_diff`.
    """
    o_norm = [_norm_token(t) for t in old_texts]
    n_norm = [_norm_token(t) for t in new_texts]
    matcher = SequenceMatcher(None, o_norm, n_norm, autojunk=False)

    old_marks: List[dict] = []
    new_marks: List[dict] = []
    changes: List[dict] = []
    counter = 0

    def _all_ph(texts):
        return bool(texts) and all(_is_placeholder(t) for t in texts)

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        old_span = old_texts[i1:i2]
        new_span = new_texts[j1:j2]
        # placeholder fill on either side => expected, not an edit
        if _all_ph(old_span) or _all_ph(new_span):
            continue
        cid = f"r{counter}"
        counter += 1
        if tag == "delete":
            kind, otype, ntype = "removed", "removed", None
        elif tag == "insert":
            kind, otype, ntype = "added", None, "added"
        else:  # replace
            kind, otype, ntype = "modified", "changed", "changed"
        for k in range(i1, i2):
            old_marks.append({"index": k, "type": otype, "change_id": cid})
        for k in range(j1, j2):
            new_marks.append({"index": k, "type": ntype, "change_id": cid})
        changes.append({
            "id": cid,
            "kind": kind,
            "old_text": " ".join(old_span),
            "new_text": " ".join(new_span),
        })
    _detect_moves_in_changes(new_marks, changes)
    return old_marks, new_marks, changes
